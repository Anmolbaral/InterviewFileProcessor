"""Focused source-integrity and parser-boundary regressions; no semantic claims."""

import hashlib
import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest import mock
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import ValidationError

from parser import (DEFAULT_MANIFEST, Corpus, Source, chunk_document, load_chunks, load_corpus, neighbors,
                    open_database, parse_manifest, parse_source, render_passages, run, save_corpus, search,
                    source_path)


MANIFEST = DEFAULT_MANIFEST


def element(tag, attributes=None, text=None):
    node = OxmlElement(tag)
    for name, value in (attributes or {}).items():
        node.set(qn(name), value)
    if text is not None:
        node.text = text
    return node


class ParserTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)

    def save(self, document, name="source.docx"):
        path = self.directory / name
        document.save(path)
        return Source(expert=1, source=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest())

    def fixture(self, *, table=False):
        document = Document()
        document.add_paragraph("")
        document.add_paragraph("Scope", style="Heading 1")
        document.add_paragraph("Expert 1  00:00:05")
        document.add_paragraph("  Alpha\tBeta\nGamma  ")
        document.add_paragraph("We measured 00:01:02 in the log.")
        heading = document.add_paragraph().add_run("Later section")
        heading.bold = heading.underline = True
        document.add_paragraph("Unlabeled answer")
        if table:
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Table-only content"
        source = self.save(document)
        manifest = self.directory / "manifest.json"
        manifest.write_text(json.dumps([source.model_dump()]), encoding="utf-8")
        return source, manifest

    def test_whitespace_locations_and_legacy_ids(self):
        source, _ = self.fixture()
        parsed = parse_source(source, self.directory)
        self.assertEqual(len(parsed.passages), 7)
        self.assertEqual([p.id for p in parsed.passages],
                         [f"E1:B{i:04d}" for i in range(1, 8)])
        self.assertEqual([p.paragraph_index for p in parsed.passages], list(range(1, 8)))
        self.assertEqual([p.citation_id for p in parsed.passages],
                         [None] + [f"E1:P{i:03d}" for i in range(1, 7)])
        self.assertEqual(parsed.passages[0].kind, "blank")
        self.assertEqual(parsed.passages[3].text, "  Alpha\tBeta\nGamma  ")
        self.assertTrue(parsed.passages[3].raw_xml.startswith("<w:p"))

    def test_only_explicit_metadata_and_heading_reset(self):
        source, _ = self.fixture()
        passages = parse_source(source, self.directory).passages
        self.assertEqual(passages[1].kind, "heading")
        self.assertEqual(passages[1].heading_basis, "style")
        self.assertEqual(passages[2].kind, "turn_marker")
        for passage in passages[2:5]:
            self.assertEqual(passage.speaker_label, "Expert 1")
            self.assertEqual(passage.timestamp_raw, "00:00:05")
            self.assertEqual(passage.turn_id, "E1:B0003")
        self.assertEqual(passages[4].kind, "text")
        self.assertEqual(passages[5].kind, "heading")
        self.assertEqual(passages[5].heading_basis, "bold_underline")
        for passage in passages[5:]:
            self.assertIsNone(passage.speaker_label)
            self.assertIsNone(passage.timestamp_raw)
            self.assertIsNone(passage.turn_id)

    def test_models_reject_wrong_types_and_unknown_output_fields(self):
        source, _ = self.fixture()
        parsed = parse_source(source, self.directory)
        for field, value in [("paragraph_index", "4"), ("unexpected", True)]:
            payload = parsed.model_dump()
            payload["passages"][3][field] = value
            with self.subTest(field=field), self.assertRaises(ValidationError):
                type(parsed).model_validate(payload)
        payload = source.model_dump()
        payload["expert"] = "1"
        with self.assertRaises(ValidationError):
            Source.model_validate(payload)

    def test_round_trip_reuse_of_unchanged_files_and_change_detection(self):
        source, manifest = self.fixture()
        first = parse_source(source, self.directory)
        second = parse_source(source, self.directory)
        self.assertEqual(first.model_dump(), second.model_dump())
        database = self.directory / "nested" / "transcripts.sqlite"
        corpus, written = run(manifest, database)
        self.assertEqual(written, ["E1"])
        connection = open_database(database, readonly=True)
        stored = Corpus.model_validate(load_corpus(connection))
        connection.close()
        self.assertEqual(stored.model_dump(), corpus.model_dump())
        # Same content hash, parser, and library: the stored rows are reused and the parser never runs.
        with mock.patch("parser.parse_source", side_effect=AssertionError("unchanged file was reparsed")):
            reused, written = run(manifest, database)
        self.assertEqual(written, [])
        self.assertEqual(reused.model_dump(), corpus.model_dump())
        before = database.read_bytes()
        self.assertEqual(run(manifest, database, check=True)[0].model_dump(), corpus.model_dump())
        self.assertEqual(database.read_bytes(), before)
        # Changed content with the old manifest hash is refused; with the new hash it is reparsed.
        path = self.directory / source.source
        document = Document(path)
        document.add_paragraph("A later answer.")
        document.save(path)
        with self.assertRaises(ValueError):
            run(manifest, database)
        manifest.write_text(json.dumps([{**source.model_dump(),
                                         "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}]), encoding="utf-8")
        corpus, written = run(manifest, database)
        self.assertEqual(written, ["E1"])
        self.assertEqual(len(corpus.documents[0].passages), 8)
        run(manifest, database, check=True)

    def test_changed_source_is_rejected(self):
        source, _ = self.fixture()
        path = self.directory / source.source
        document = Document(path)
        document.add_paragraph("A changed source must not inherit the old fingerprint.")
        document.save(path)
        with self.assertRaises(ValueError):
            parse_source(source, self.directory)

    def test_table_reports_partial_coverage(self):
        source, _ = self.fixture(table=True)
        parsed = parse_source(source, self.directory)
        self.assertEqual(parsed.coverage, "partial")
        self.assertTrue(any("tabl" in warning.lower() for warning in parsed.warnings))

    def test_corrupt_docx_is_rejected(self):
        path = self.directory / "corrupt.docx"
        path.write_bytes(b"This is not an Office ZIP package.")
        source = Source(expert=1, source=path.name,
                        sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        with self.assertRaises(ValueError):
            parse_source(source, self.directory)

    def test_nested_runs_are_kept_and_body_markers_are_not_structures(self):
        document = Document()
        paragraph = document.add_paragraph("Plain run. ")
        document.add_paragraph("No speaker labels anywhere in this file.")
        body = document.element.body
        body.insert(0, element("w:bookmarkStart", {"w:id": "0", "w:name": "_GoBack"}))
        body.insert(1, element("w:bookmarkEnd", {"w:id": "0"}))
        parsed = parse_source(self.save(document), self.directory)
        self.assertEqual(parsed.coverage, "supported_text")
        self.assertEqual(len(parsed.warnings), 1)
        self.assertIn("no speaker turn markers", parsed.warnings[0])
        # A run inside a wrapper element that python-docx's paragraph text skips must still be stored.
        wrapper = element("w:smartTag", {"w:element": "PersonName"})
        run_ = element("w:r")
        run_.append(element("w:t", text="INSIDE_SMARTTAG"))
        wrapper.append(run_)
        paragraph._p.append(wrapper)
        parsed = parse_source(self.save(document), self.directory)
        self.assertEqual(parsed.passages[0].text, "Plain run. INSIDE_SMARTTAG")
        self.assertEqual(parsed.coverage, "partial")
        self.assertTrue(any(w.startswith("E1:B0001:") and "run content" in w for w in parsed.warnings))

    def test_tampered_rows_and_dangling_references_are_rejected(self):
        _, manifest = self.fixture()
        database = self.directory / "transcripts.sqlite"
        for column, value in [("citation_id", "E9:P999"), ("text", "Tampered text")]:
            database.unlink(missing_ok=True)
            run(manifest, database)
            connection = open_database(database)
            with connection:
                connection.execute(f"UPDATE passages SET {column} = ? WHERE id = 'E1:B0004'", (value,))
            connection.close()
            before = database.read_bytes()
            with self.subTest(column=column), self.assertRaises(ValueError):
                run(manifest, database, check=True)
            # The write path refuses to reuse rows that fail validation instead of trusting them.
            with self.subTest(column=column, mode="write"), self.assertRaises(ValueError):
                run(manifest, database)
            self.assertEqual(database.read_bytes(), before)
        connection = open_database(database)
        with self.assertRaises(sqlite3.IntegrityError), connection:
            connection.execute("UPDATE passages SET section_id = 'E1:B9999' WHERE id = 'E1:B0004'")
        connection.close()

    def test_grouping_edge_cases_and_derived_integrity(self):
        _, manifest = self.fixture()  # a two-paragraph expert turn, then unattributed text after a heading
        database = self.directory / "transcripts.sqlite"
        corpus, _ = run(manifest, database)
        connection = open_database(database)
        chunks = load_chunks(connection)
        self.assertEqual(
            [(c["id"], c["level"], c["passage_ids"], c["speaker_labels"], c["section_id"], c["header"],
              c["header_passage_ids"]) for c in chunks],
            [("E1:X0001", "exchange", ["E1:B0004", "E1:B0005"], ["Expert 1"], "E1:B0002", "E1 | Scope", ["E1:B0002"]),
             ("E1:X0002", "unattributed", ["E1:B0007"], [], "E1:B0006", "E1 | Later section", ["E1:B0006"])])
        passages = {p.id: p for p in corpus.documents[0].passages}
        for chunk in chunks:
            self.assertEqual(chunk["text"], "\n".join(passages[i].text for i in chunk["passage_ids"]))
        self.assertEqual([c["id"] for c in neighbors(connection, "E1:X0001")], ["E1:X0001"])  # next is another section
        self.assertEqual(neighbors(connection, "E1:X9999"), [])
        rendered = render_passages(connection, ["E1:B0004", "E1:B0005"],
                                   expected_fingerprints={"E1": corpus.documents[0].extraction_sha256})
        self.assertEqual([(r["citation_id"], r["speaker_label"], r["timestamp_raw"]) for r in rendered],
                         [("E1:P003", "Expert 1", "00:00:05"), ("E1:P004", "Expert 1", "00:00:05")])
        self.assertEqual(rendered[0]["text"], "  Alpha\tBeta\nGamma  ")
        self.assertEqual([h["id"] for h in search(connection, "Gamma")], ["E1:X0001"])
        self.assertEqual([h["id"] for h in search(connection, "Later section?")], ["E1:X0002"])  # header, fallback
        with self.assertRaises(sqlite3.IntegrityError), connection:
            connection.execute("INSERT INTO chunk_passages VALUES ('E1:X0001', 'E1:B9999', 9)")
        connection.close()
        for statement in ("UPDATE chunks SET text = 'Tampered' WHERE id = 'E1:X0001'",
                          "UPDATE chunks_fts SET text = 'Tampered' WHERE chunk_id = 'E1:X0001'"):
            run(manifest, database)  # A normal run rebuilds every derived row.
            connection = open_database(database)
            with connection:
                connection.execute(statement)
            connection.close()
            before = database.read_bytes()
            with self.subTest(statement=statement), self.assertRaises(ValueError):
                run(manifest, database, check=True)
            self.assertEqual(database.read_bytes(), before)
        run(manifest, database)
        run(manifest, database, check=True)
        connection = open_database(database)
        with connection:
            connection.execute("DELETE FROM documents")
        self.assertEqual([connection.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
                          for t in ("chunks", "chunk_passages")], [0, 0])
        connection.close()

    def test_render_passages_fails_on_missing_duplicate_or_stale_evidence(self):
        _, manifest = self.fixture()
        database = self.directory / "transcripts.sqlite"
        corpus, _ = run(manifest, database)
        document = corpus.documents[0]
        expected = {"E1": document.extraction_sha256}
        connection = open_database(database, readonly=True)
        self.addCleanup(connection.close)
        rows = render_passages(connection, ["E1:B0007", "E1:B0005", "E1:B0003", "E1:B0004"],
                               expected_fingerprints=expected)  # shuffled input spanning two sections
        self.assertEqual([r["id"] for r in rows], ["E1:B0003", "E1:B0004", "E1:B0005", "E1:B0007"])
        self.assertEqual([(r["citation_id"], r["speaker_label"], r["timestamp_raw"]) for r in rows],
                         [("E1:P002", "Expert 1", "00:00:05"), ("E1:P003", "Expert 1", "00:00:05"),
                          ("E1:P004", "Expert 1", "00:00:05"), ("E1:P006", None, None)])
        self.assertEqual(rows[1]["text"], "  Alpha\tBeta\nGamma  ")
        self.assertEqual({(r["document_id"], r["source_sha256"], r["extraction_sha256"]) for r in rows},
                         {("E1", document.source_sha256, document.extraction_sha256)})
        self.assertEqual(render_passages(connection, []), [])
        for ids in (["E1:B0004", "E1:B9999"], ["E1:B0004", "E1:B0004"]):
            with self.subTest(ids=ids), self.assertRaises(ValueError):
                render_passages(connection, ids, expected_fingerprints=expected)
        with self.assertRaisesRegex(ValueError, "expected.*fingerprint"):
            render_passages(connection, ["E1:B0004"])
        for invalid in (None, {}, {"E2": document.extraction_sha256}, {"E1": None}, {"E1": "0" * 64}):
            with self.subTest(expected=invalid), self.assertRaisesRegex(ValueError, "expected.*fingerprint"):
                render_passages(connection, ["E1:B0004"], expected_fingerprints=invalid)

    def test_search_cli_checks_integrity_before_returning_evidence(self):
        source, manifest = self.fixture()
        database = self.directory / "transcripts.sqlite"
        run(manifest, database)
        pristine = database.read_bytes()

        def invoke(query="Gamma"):
            return subprocess.run([sys.executable, str(Path(__file__).with_name("parser.py")),
                                   "--manifest", str(manifest), "--database", str(database), "--search", query],
                                  capture_output=True, text=True, check=False)

        result = invoke()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("E1:P003 Expert 1", result.stdout)
        self.assertEqual(database.read_bytes(), pristine)
        for statement in ("UPDATE passages SET text = 'Tampered' WHERE id = 'E1:B0004'",
                          "UPDATE chunks_fts SET text = 'Tampered' WHERE chunk_id = 'E1:X0001'"):
            database.write_bytes(pristine)
            connection = open_database(database)
            with connection:
                connection.execute(statement)
            connection.close()
            before = database.read_bytes()
            for query in ("Gamma", "nonexistent"):
                with self.subTest(statement=statement, query=query):
                    result = invoke(query)
                    self.assertEqual(result.returncode, 1, result.stdout)
                    self.assertIn("Parser failed:", result.stderr)
                    self.assertEqual(result.stdout, "")
                    self.assertEqual(database.read_bytes(), before)
        database.write_bytes(pristine)
        path = self.directory / source.source
        document = Document(path)
        document.add_paragraph("A source change after the database was built.")
        document.save(path)
        result = invoke()
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("Parser failed:", result.stderr)
        self.assertEqual(result.stdout, "")
        self.assertEqual(database.read_bytes(), pristine)

    def test_failure_preserves_existing_database(self):
        _, manifest = self.fixture()
        database = self.directory / "transcripts.sqlite"
        run(manifest, database)
        before = database.read_bytes()
        entries = json.loads(manifest.read_text(encoding="utf-8"))
        entries.append({"expert": 2, "source": "missing.docx", "sha256": "0" * 64})
        manifest.write_text(json.dumps(entries), encoding="utf-8")
        with self.assertRaises((ValueError, OSError)):
            run(manifest, database)
        self.assertEqual(database.read_bytes(), before)

    def test_database_cannot_replace_manifest_or_source(self):
        source, manifest = self.fixture()
        for database in (manifest, self.directory / source.source):
            before = database.read_bytes()
            with self.subTest(database=database.name), self.assertRaises(ValueError):
                run(manifest, database)
            self.assertEqual(database.read_bytes(), before)


@unittest.skipUnless(MANIFEST.is_file(), "Transcript manifest (inputs/manifest.json) is missing; synthetic tests still run.")
class ProvidedTranscriptTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.corpus = parse_manifest(MANIFEST)

    def test_counts_and_read_only_source_hashes(self):
        self.assertEqual(self.corpus.schema_version, "1")
        self.assertEqual([document.id for document in self.corpus.documents], ["E1", "E2", "E3"])
        self.assertEqual([len(d.passages) for d in self.corpus.documents], [494, 169, 157])
        self.assertEqual([sum(p.citation_id is not None for p in d.passages)
                          for d in self.corpus.documents], [326, 169, 157])
        for entry, document in zip(json.loads(MANIFEST.read_text()), self.corpus.documents):
            self.assertEqual(hashlib.sha256(source_path(entry["source"], MANIFEST.parent).read_bytes()).hexdigest(),
                             entry["sha256"])
            self.assertEqual(document.source_sha256, entry["sha256"])
            # Independent run-level XML comparison catches dropped/inserted whitespace.
            w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            with zipfile.ZipFile(source_path(entry["source"], MANIFEST.parent)) as archive:
                body = ET.fromstring(archive.read("word/document.xml")).find(w + "body")
            originals = body.findall(w + "p")
            self.assertEqual(len(originals), len(document.passages))
            for original, passage in zip(originals, document.passages):
                expected = "".join(
                    (node.text or "") if node.tag == w + "t" else
                    "\t" if node.tag == w + "tab" else
                    "\n" if node.tag == w + "cr" or (
                        node.tag == w + "br" and node.get(w + "type", "textWrapping") == "textWrapping"
                    ) else ""
                    for run in original.iter(w + "r") for node in run
                )
                self.assertEqual(passage.text, expected, passage.id)

    def test_exchange_chunks_search_and_context_on_supplied_files(self):
        dump = self.corpus.model_dump()
        expected = {d.id: d.extraction_sha256 for d in self.corpus.documents}
        with tempfile.TemporaryDirectory() as directory:
            connection = open_database(Path(directory) / "transcripts.sqlite")
            save_corpus(connection, dump)
            chunks = load_chunks(connection)
            per_document = [[c for c in chunks if c["document_id"] == d] for d in ("E1", "E2", "E3")]
            self.assertEqual([sum(c["level"] == "exchange" for c in cs) for cs in per_document], [80, 40, 37])
            self.assertEqual([sum(c["speaker_labels"][:1] == ["AI Interviewer"] for c in cs) for cs in per_document],
                             [79, 39, 36])
            self.assertEqual([sum(c["level"] == "unattributed" for c in cs) for cs in per_document], [1, 1, 1])
            text_ids = sorted(p["id"] for d in dump["documents"] for p in d["passages"] if p["kind"] == "text")
            self.assertEqual(sorted(i for c in chunks for i in c["passage_ids"]), text_ids)  # exactly one chunk each
            self.assertEqual(chunks, [c for d in dump["documents"] for c in chunk_document(d)])
            price_id = next(p["id"] for p in dump["documents"][2]["passages"] if p["citation_id"] == "E3:P098")
            chunk = next(c for c in chunks if price_id in c["passage_ids"])
            self.assertEqual(chunk["header"], "E3 | Expert 3 Head of IT Operations at Solara Renewables "
                                              "(2020 - present) | Cost & Total Cost of Ownership")
            self.assertEqual(chunk["header_passage_ids"], ["E3:B0001", "E3:B0002", chunk["section_id"]])
            self.assertEqual(search(connection, '"forty dollars per agent"', document_id="E3")[0]["id"], chunk["id"])
            self.assertIn(chunk["id"], [h["id"] for h in search(connection, "Solara pricing", document_id="E3",
                                                                   limit=50)])  # via the header; rank not asserted
            rendered = render_passages(connection, chunk["passage_ids"], expected_fingerprints=expected)
            self.assertEqual([(r["speaker_label"], r["citation_id"]) for r in rendered],
                             [("AI Interviewer", "E3:P096"), ("Expert 3", "E3:P098")])
            # Company context and later corrections remain accessible across sections and interviews.
            citations = {p.citation_id: p for d in self.corpus.documents for p in d.passages}
            ordered = ["E1:P016", "E1:P021", "E1:P049", "E2:P089", "E2:P139"]
            ids = [citations[c].id for c in reversed(ordered)]
            self.assertNotEqual(citations["E2:P089"].section_id, citations["E2:P139"].section_id)
            with self.assertRaisesRegex(ValueError, "E2.*expected.*fingerprint"):
                render_passages(connection, ids, expected_fingerprints={"E1": expected["E1"]})
            rendered = render_passages(connection, ids, expected_fingerprints=expected)
            self.assertEqual([r["citation_id"] for r in rendered], ordered)
            self.assertEqual([r["text"] for r in rendered], [citations[c].text for c in ordered])
            around = neighbors(connection, chunk["id"])
            self.assertEqual([c["id"] for c in around],
                             [c["id"] for c in per_document[2] if c["section_id"] == chunk["section_id"]
                              and abs(c["ordinal"] - chunk["ordinal"]) <= 1])
            self.assertIn(chunk["id"], [c["id"] for c in around])
            connection.close()

    def test_golden_citations_newlines_and_no_formatting_tabs(self):
        citations = {p.citation_id: p for d in self.corpus.documents for p in d.passages
                     if p.citation_id is not None}
        self.assertIn("previous stop, Thermo Fisher", citations["E1:P021"].text)
        self.assertIn("roughly forty dollars per agent per month", citations["E3:P098"].text)
        expert1_text = "".join(p.text for p in self.corpus.documents[0].passages)
        self.assertEqual(expert1_text.count("\n"), 2)
        self.assertNotIn("\t", expert1_text)
        self.assertTrue(citations["E1:P001"].text.startswith("\nExpert 1\n"))
        ancillary = {part.part_name: part for part in self.corpus.documents[0].ancillary_parts}
        self.assertIn("SYNQUERY - FULL TRANSCRIPT",
                      "\n".join(ancillary["word/header1.xml"].paragraphs))
        self.assertNotIn("SYNQUERY - FULL TRANSCRIPT", expert1_text)
        self.assertTrue(ancillary["word/footer1.xml"].field_instructions)
        for document in self.corpus.documents:
            # Both text readers agree on every paragraph of both transcript templates.
            self.assertFalse(any("run content" in warning for warning in document.warnings), document.id)
            self.assertEqual(document.coverage, "supported_text")
            for passage in document.passages:
                if passage.kind == "heading":
                    self.assertIsNone(passage.speaker_label)
                    self.assertIsNone(passage.turn_id)
                    self.assertIsNone(passage.timestamp_raw)


if __name__ == "__main__":
    unittest.main()
