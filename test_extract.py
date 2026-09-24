"""Contract, schedule, validation, storage, and bounded evidence-guard checks. No model calls."""

import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import threading
import typing
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from pydantic import ValidationError

import extract
from extract import (PROMPT_VERSION, BatchOutcome, CompanyContext, ExtractionRun, Finding, ModelReply, Quantity,
                     SourceVersion, attribution_flags, build_batches, check_findings, company_anchors, extract_batches,
                     calibrate, judge_cases, link_batches, model_batches, run_identity,
                     load_env, load_records, load_runs, normalize_quantity, open_findings, realize, render_batch,
                     review, save_run, user_message, validate_record)
from extract import DEFAULT_FINDINGS, Proposal, evaluate
from parser import DEFAULT_DATABASE, DEFAULT_MANIFEST, Source, open_database, run

MANIFEST = DEFAULT_MANIFEST
FP = "0" * 64


def write_fixture(directory: Path, experts: tuple[int, ...] = (1,)) -> tuple[Path, Path]:
    """Two sections behind a profile line that precedes the first heading, the shape of the supplied files.
    One interview per expert number, identical apart from the speaker label; returns the first path and the manifest."""
    sources, paths = [], []
    for expert in experts:
        document = Document()
        document.add_paragraph(f"Expert {expert}")
        document.add_paragraph("Introduction", style="Heading 1")
        document.add_paragraph("AI Interviewer  00:00:01")
        document.add_paragraph("Tell me about your role.")
        document.add_paragraph(f"Expert {expert}  00:00:05")
        document.add_paragraph("I lead IT at Acme, previously at Globex.")
        document.add_paragraph("Cost", style="Heading 1")
        document.add_paragraph("AI Interviewer  00:01:00")
        document.add_paragraph("What do you pay per agent?")
        document.add_paragraph(f"Expert {expert}  00:01:10")
        document.add_paragraph("Roughly forty dollars per agent per month.")
        document.add_paragraph("AI Interviewer  00:02:00")
        document.add_paragraph("Any surprises?")
        document.add_paragraph(f"Expert {expert}  00:02:05")
        document.add_paragraph("Yes, but that was testing access, not extra environments.")
        path = directory / f"source{expert}.docx"
        document.save(path)
        paths.append(path)
        sources.append(Source(expert=expert, source=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    manifest = directory / "manifest.json"
    manifest.write_text(json.dumps([source.model_dump() for source in sources]), encoding="utf-8")
    return paths[0], manifest


def write_history_fixture(directory: Path) -> Path:
    """An introduction naming both employers, a section about the former one, then a cost section that names neither:
    the shape of Expert 1's Thermo Fisher account. Returns the manifest."""
    document = Document()
    document.add_paragraph("Expert 1")
    for heading, turns in (("Introduction", ("Tell me about your role.", "I lead IT at Acme, previously at Globex.")),
                           ("Globex", ("What did Globex run?", "We ran ServiceNow there for years.")),
                           ("Cost", ("How did total cost compare?", "Roughly double the alternatives."))):
        document.add_paragraph(heading, style="Heading 1")
        document.add_paragraph("AI Interviewer  00:00:01")
        document.add_paragraph(turns[0])
        document.add_paragraph("Expert 1  00:00:05")
        document.add_paragraph(turns[1])
    path = directory / "history.docx"
    document.save(path)
    manifest = directory / "manifest.json"
    manifest.write_text(json.dumps([Source(expert=1, source=path.name, sha256=hashlib.sha256(path.read_bytes())
                                           .hexdigest()).model_dump()]), encoding="utf-8")
    return manifest


class FakeModel:
    """Stands in for the provider call: replies in order, raises when given an exception, records every prompt."""

    def __init__(self, replies):
        self.replies, self.calls = list(replies), []

    def __call__(self, system, user):
        self.calls.append((system, user))
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        return ModelReply(text=reply, model="fake-model", finish="stop")


class LaneModel:
    """Replies from the prompt it receives: the introduction batch proposes the company, the cost batch prices it
    and reuses the context the prompt lists, other batches propose nothing. Records which thread each call ran on."""

    def __init__(self):
        self.calls, self.on_main_thread = [], []

    def __call__(self, system, user):
        self.calls.append(user)
        self.on_main_thread.append(threading.current_thread() is threading.main_thread())
        document = re.search(r"Existing company contexts for interview (E\d+)", user).group(1)
        if f"[{document}:P011]" in user:
            listed = re.search(rf"- ({document}:C\d+):", user)
            contexts = [] if listed else [dict(key="acme", company_name="Acme", employment="current",
                                               supporting=[f"{document}:P006"])]
            finding = dict(context_ref=listed.group(1) if listed else "acme", kind="pricing", vendor="Acme Desk",
                           statement="The expert reports roughly forty dollars per agent per month.",
                           evidence_type="expert_estimate",
                           quantity=dict(raw="Roughly forty dollars per agent per month.", value=40.0, approximate=True,
                                         unit="per agent", period="month", currency="dollars"),
                           supporting=[f"{document}:P011"], qualifying=[f"{document}:P009"], context=[f"{document}:P006"])
            reply = dict(contexts=contexts, findings=[finding])
        elif f"[{document}:P006]" in user:
            reply = dict(contexts=[dict(key="acme", company_name="Acme", employment="current",
                                        supporting=[f"{document}:P006"])], findings=[])
        else:
            reply = dict(contexts=[], findings=[])
        return ModelReply(text=json.dumps(reply), model="fake-model", finish="stop")


class ExtractTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.source, self.manifest = write_fixture(self.directory)
        self.database = self.directory / "transcripts.sqlite"
        corpus, _ = run(self.manifest, self.database)
        self.document = corpus.documents[0]
        self.expected = {"E1": self.document.extraction_sha256}
        self.connection = open_database(self.database, readonly=True)
        self.addCleanup(self.connection.close)

    def version(self):
        return dict(source_sha256=self.document.source_sha256, extraction_sha256=self.document.extraction_sha256)

    def context(self, **overrides):
        base = dict(id="E1:C01", document_id="E1", company_name="Acme", employment="current",
                    supporting=["E1:B0006"], origin="analyst", **self.version())
        return CompanyContext(**{**base, **overrides})

    def finding(self, **overrides):
        base = dict(id="E1:F001", document_id="E1", context_id="E1:C01", kind="pricing", vendor="Acme Desk",
                    statement="The expert reports roughly forty dollars per agent per month.",
                    evidence_type="expert_estimate",
                    quantity=Quantity(raw="Roughly forty dollars per agent per month.", value=40.0, approximate=True,
                                      unit="per agent", period="month", currency="dollars"),
                    supporting=["E1:B0011"], qualifying=["E1:B0009"], context=["E1:B0006"], origin="analyst",
                    **self.version())
        return Finding(**{**base, **overrides})

    def run_record(self, **overrides):
        base = dict(id="run-1", model="none", prompt_version="analyst-seed",
                    source_versions={"E1": SourceVersion(**self.version())})
        return ExtractionRun(**{**base, **overrides})

    def test_batches_cover_every_text_passage_once_with_navigation_and_context(self):
        batches = build_batches(self.connection)
        self.assertEqual([(b.id, b.heading, b.passage_ids, b.context_ids, b.chunk_ids) for b in batches], [
            ("E1:S00", None, ["E1:B0001"], [], ["E1:X0001"]),
            ("E1:S01", "Introduction", ["E1:B0004", "E1:B0006"], ["E1:B0001"], ["E1:X0002"]),
            ("E1:S02", "Cost", ["E1:B0009", "E1:B0011", "E1:B0013", "E1:B0015"],
             ["E1:B0001", "E1:B0004", "E1:B0006"], ["E1:X0003", "E1:X0004"]),
        ])
        self.assertEqual([b.section_id for b in batches], [None, "E1:B0002", "E1:B0007"])
        text_ids = [row[0] for row in self.connection.execute("SELECT id FROM passages WHERE kind = 'text'")]
        self.assertEqual(sorted(i for b in batches for i in b.passage_ids), sorted(text_ids))
        chunk_ids = [row[0] for row in self.connection.execute("SELECT id FROM chunks")]
        self.assertEqual(sorted(i for b in batches for i in b.chunk_ids), sorted(chunk_ids))

    def test_batch_input_is_canonical_attributed_and_version_bound(self):
        batch = build_batches(self.connection)[2]
        text = render_batch(self.connection, batch, expected_fingerprints=self.expected)
        self.assertIn("Section heading (navigation only, not evidence): Cost", text)
        self.assertIn("[E1:P006] Expert 1 00:00:05: I lead IT at Acme, previously at Globex.", text)
        self.assertIn("[E1:P009] AI Interviewer 00:01:00: What do you pay per agent?", text)
        before, after = text.split("Section passages (primary evidence):")
        self.assertIn("Introductory context", before)
        self.assertIn("[E1:P006]", before)
        self.assertNotIn("[E1:P011]", before)
        self.assertIn("[E1:P011] Expert 1 00:01:10: Roughly forty dollars per agent per month.", after)
        self.assertNotIn("[E1:P006]", after)
        for invalid in ({"E1": "0" * 64}, {}):
            with self.subTest(expected=invalid), self.assertRaisesRegex(ValueError, "fingerprint"):
                render_batch(self.connection, batch, expected_fingerprints=invalid)

    def test_records_enforce_the_contract(self):
        Quantity(raw="north of one-fifty", low=150.0)  # an open bound keeps the source's vagueness
        self.context()
        self.finding()
        self.finding(quantity=None, statement="The overrun came mostly from extra agent seats.")  # cost without a figure
        self.finding(kind="vendor_relationship", relationship="deployed", vendor=None,
                     statement="The current employer runs a custom ITSM.")  # an unnamed or in-house system
        self.finding(kind="rating", quantity=Quantity(raw="9 out of 10", value=9.0, unit="out of 10"),
                     statement="The expert rates renewal likelihood 9 out of 10.")  # scores have a kind of their own
        self.run_record(batches=[BatchOutcome(batch_id="E1:S02", status="findings", input_passage_ids=["E1:B0011"],
                                              input_fingerprint=FP, finding_ids=["E1:F001"])])
        rejected = [
            ("range reversed", lambda: Quantity(raw="x", low=55.0, high=50.0)),
            ("value outside bounds", lambda: Quantity(raw="x", value=100.0, low=50.0, high=55.0)),
            ("vendor relationship without relationship", lambda: self.finding(kind="vendor_relationship")),
            ("no supporting passage", lambda: self.finding(supporting=[])),
            ("another interview's passage", lambda: self.finding(context=["E2:B0006"])),
            ("citation ID instead of physical ID", lambda: self.finding(supporting=["E1:P011"])),
            ("one passage in two roles", lambda: self.finding(context=["E1:B0011"])),
            ("stored supporting/qualifying overlap", lambda: self.finding(
                qualifying=["E1:B0011"], qualifications=["The expert estimates this amount."])),
            ("unknown evidence type", lambda: self.finding(evidence_type="fact")),
            ("unknown field", lambda: self.finding(confidence=0.9)),
            ("empty statement", lambda: self.finding(statement="")),
            ("unknown employment", lambda: self.context(employment="prior")),
            ("failed batch without error", lambda: BatchOutcome(batch_id="E1:S02", status="failed",
                                                                 input_passage_ids=[], input_fingerprint=FP)),
            ("no-findings batch with findings", lambda: BatchOutcome(batch_id="E1:S02", status="no_findings", input_passage_ids=[],
                                                                     input_fingerprint=FP, finding_ids=["E1:F001"])),
            ("duplicate batch outcomes", lambda: self.run_record(batches=[
                BatchOutcome(batch_id="E1:S02", status="no_findings", input_passage_ids=[], input_fingerprint=FP),
                BatchOutcome(batch_id="E1:S02", status="no_findings", input_passage_ids=[], input_fingerprint=FP)])),
        ]
        for label, build in rejected:
            with self.subTest(label), self.assertRaises(ValidationError):
                build()

    def test_validation_rejects_unsupplied_nontext_unknown_or_stale_evidence(self):
        batch = build_batches(self.connection)[2]
        supplied = set(batch.passage_ids) | set(batch.context_ids)
        rows = validate_record(self.connection, self.finding(), supplied, expected_fingerprints=self.expected)
        self.assertEqual([r["citation_id"] for r in rows], ["E1:P006", "E1:P009", "E1:P011"])  # source order
        self.assertEqual(rows[2]["text"], "Roughly forty dollars per agent per month.")
        cases = [
            ("not supplied to the batch", self.finding(context=["E1:B0004"]), supplied - {"E1:B0004"}, "E1:B0004"),
            ("turn marker", self.finding(supporting=["E1:B0010"]), supplied | {"E1:B0010"}, "text passage"),
            ("heading", self.finding(supporting=["E1:B0007"]), supplied | {"E1:B0007"}, "text passage"),
            ("unknown passage", self.finding(supporting=["E1:B9999"]), supplied | {"E1:B9999"}, "unknown"),
            ("bound to another version", self.finding(extraction_sha256="0" * 64), supplied, "fingerprint"),
        ]
        for label, record, ids, message in cases:
            with self.subTest(label), self.assertRaisesRegex(ValueError, f"(?s)E1:F001.*{message}"):
                validate_record(self.connection, record, ids, expected_fingerprints=self.expected)

    def test_store_never_overwrites_and_check_flags_stale_records_without_relabeling(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        context, finding = self.context(), self.finding(review_state="reviewed")
        save_run(store, self.run_record(), [context], [finding])
        self.assertEqual(load_records(store, CompanyContext), [context])
        self.assertEqual(load_records(store, Finding), [finding])
        self.assertEqual(check_findings(store, self.connection), (1, 1))
        with self.assertRaisesRegex(ValueError, "E1:F001"):
            save_run(store, self.run_record(), [], [self.finding(statement="Rewritten")])
        with self.assertRaisesRegex(ValueError, "E1:F002"):
            save_run(store, self.run_record(), [], [self.finding(id="E1:F002", extraction_sha256="0" * 64)])
        with self.assertRaisesRegex(ValueError, "E1:F003"):
            save_run(store, self.run_record(), [], [self.finding(id="E1:F003", context_id="E1:C99")])
        self.assertEqual(load_records(store, Finding), [finding])
        # The source changes and the transcripts database is rebuilt: the reviewed finding is stale, not relabeled.
        self.connection.close()
        document = Document(self.source)
        document.add_paragraph("A later addition.")
        document.save(self.source)
        entries = json.loads(self.manifest.read_text(encoding="utf-8"))
        entries[0]["sha256"] = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.manifest.write_text(json.dumps(entries), encoding="utf-8")
        run(self.manifest, self.database)
        transcripts = open_database(self.database, readonly=True)
        self.addCleanup(transcripts.close)
        with self.assertRaisesRegex(ValueError, "(?s)E1:C01.*fingerprint.*E1:F001.*fingerprint"):
            check_findings(store, transcripts)
        self.assertEqual(load_records(store, Finding), [finding])

    def proposal(self, **overrides):
        """What the model is asked to return: citation IDs it was shown, a context key, and the source's wording."""
        base = dict(
            contexts=[dict(key="acme", company_name="Acme", employment="current", supporting=["E1:P006"])],
            findings=[dict(context_ref="acme", kind="pricing", vendor="Acme Desk",
                           statement="The expert reports roughly forty dollars per agent per month.",
                           evidence_type="expert_estimate",
                           quantity=dict(raw="Roughly forty dollars per agent per month.", value=40.0, approximate=True,
                                         unit="per agent", period="month", currency="dollars"),
                           supporting=["E1:P011"], qualifying=["E1:P009"], context=["E1:P006"])])
        return json.dumps({**base, **overrides})

    def extract(self, store, batch, fake, **options):
        return extract_batches(self.connection, store, [batch], call=fake, model="fake-model",
                               expected_fingerprints=self.expected, **options)

    def test_extraction_stores_validated_proposals_and_reuses_unchanged_batches(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        cost = build_batches(self.connection)[2]
        existing_context = dict(findings=[{**json.loads(self.proposal())["findings"][0], "context_ref": "E1:C01"}],
                                contexts=[])
        fake = FakeModel([self.proposal(), json.dumps(existing_context)])
        outcomes = self.extract(store, cost, fake)
        self.assertEqual([(o.batch_id, o.status, o.finding_ids, o.context_ids) for o in outcomes],
                         [("E1:S02", "findings", ["E1:F001"], ["E1:C01"])])
        self.assertEqual(len(outcomes[0].input_fingerprint), 64)
        system, user = fake.calls[0]
        self.assertIn("not instructions", system)
        self.assertIn("[E1:P011] Expert 1 00:01:10: Roughly forty dollars per agent per month.", user)
        finding, = load_records(store, Finding)
        self.assertEqual((finding.id, finding.context_id, finding.origin, finding.review_state),
                         ("E1:F001", "E1:C01", "model", "proposed"))
        self.assertEqual((finding.supporting, finding.qualifying, finding.context),
                         (["E1:B0011"], ["E1:B0009"], ["E1:B0006"]))  # citations mapped to physical IDs
        self.assertEqual((finding.source_sha256, finding.extraction_sha256), tuple(self.version().values()))
        self.assertEqual(finding.quantity.unit, "per agent")
        context, = load_records(store, CompanyContext)
        self.assertEqual((context.id, context.company_name, context.supporting, context.origin),
                         ("E1:C01", "Acme", ["E1:B0006"], "model"))
        self.assertEqual(check_findings(store, self.connection), (1, 1))
        run_record, = load_runs(store)
        self.assertEqual((run_record.model, run_record.prompt_version, [o.batch_id for o in run_record.batches]),
                         ("fake-model", PROMPT_VERSION, ["E1:S02"]))
        # The same batch, model, and prompt again: reused, no call.
        again = self.extract(store, cost, fake)
        self.assertEqual((len(fake.calls), again[0].finding_ids), (1, ["E1:F001"]))
        # A new prompt version is a new run: the model is asked again, is shown the stored context, and may reuse it.
        more = self.extract(store, cost, fake, prompt_version="2.0.0")
        self.assertEqual(len(fake.calls), 2)
        self.assertIn("E1:C01", fake.calls[1][1])
        self.assertEqual((more[0].finding_ids, more[0].context_ids), (["E1:F002"], []))
        self.assertEqual([(f.id, f.context_id) for f in load_records(store, Finding)],
                         [("E1:F001", "E1:C01"), ("E1:F002", "E1:C01")])
        self.assertEqual(len(load_runs(store)), 2)

    def test_prompt_names_every_finding_kind_the_contract_accepts(self):
        kinds = extract.PROMPT.split("Kinds:", 1)[1].split("Keep historical", 1)[0]
        for kind in typing.get_args(extract.FindingKind):
            self.assertIn(kind, kinds)

    def test_links_read_the_interviews_stored_findings_and_cite_only_their_passages(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        self.extract(store, build_batches(self.connection)[2], FakeModel([self.proposal()]))
        run_id = run_identity(self.connection, "fake-model")
        link, = link_batches(self.connection, store, run_id)
        self.assertEqual((link.id, link.passage_ids, link.context_ids), ("E1:L01", ["E1:B0009", "E1:B0011"],
                                                                         ["E1:B0006"]))
        linked = dict(context_ref="E1:C01", kind="pricing", evidence_type="firsthand_report",
                      statement="The forty-dollar answer and the question it answers are one account.",
                      supporting=["E1:P011"], qualifying=["E1:P009"], qualifications=["Nothing is reconciled."])
        stray = {**linked, "supporting": ["E1:P015"], "qualifying": []}  # no stored finding cites P015
        company = dict(key="globex", company_name="Globex", employment="historical", supporting=["E1:P006"])
        fake = FakeModel([json.dumps(dict(contexts=[company], findings=[linked, stray]))])
        outcome, = self.extract(store, link, fake)
        system, user = fake.calls[0]
        self.assertEqual(system, extract.LINK_PROMPT)
        self.assertIn("E1:F001", user)
        self.assertIn("[E1:P011] Expert 1 00:01:10: Roughly forty dollars per agent per month.", user)
        self.assertEqual((outcome.status, outcome.finding_ids, outcome.context_ids), ("findings", ["E1:F002"], []))
        self.assertEqual(len(outcome.rejected), 2)  # links propose no companies and cite only stored findings' passages
        made = next(f for f in load_records(store, Finding) if f.id == "E1:F002")
        self.assertEqual((made.supporting, made.qualifying), (["E1:B0011"], ["E1:B0009"]))
        # The link's own finding stays out of the next link input, so an unchanged rerun is reused without a call.
        again, = self.extract(store, link_batches(self.connection, store, run_id)[0], fake)
        self.assertEqual((len(fake.calls), again.finding_ids), (1, ["E1:F002"]))

    def test_a_rerun_link_never_reads_earlier_link_findings_or_other_runs(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        save_run(store, self.run_record(id="analyst"), [self.context()], [])  # another run's company
        other = self.finding(id="E1:F001", statement="An analyst's finding from another run.")
        save_run(store, self.run_record(id="analyst"), [], [other])
        cost = build_batches(self.connection)[2]
        section = dict(contexts=[], findings=[{**json.loads(self.proposal())["findings"][0], "context_ref": "E1:C01"}])
        self.extract(store, cost, FakeModel([json.dumps(section)]))  # E1:F002
        run_id = run_identity(self.connection, "fake-model")
        linked = dict(context_ref="E1:C01", kind="pricing", evidence_type="firsthand_report",
                      statement="A link.", supporting=["E1:P011"], qualifying=["E1:P009"],
                      qualifications=["Nothing is reconciled."])
        first = FakeModel([json.dumps(dict(contexts=[], findings=[linked]))])
        self.extract(store, link_batches(self.connection, store, run_id)[0], first)  # E1:F003
        self.assertNotIn("E1:F001", first.calls[0][1])  # another run's finding is not this run's to link
        # Another section adds a finding, so the link reruns; it must not read its own earlier finding back.
        role = dict(context_ref="E1:C01", kind="company_scale", evidence_type="firsthand_report",
                    statement="The expert leads IT at Acme.", supporting=["E1:P006"])
        self.extract(store, build_batches(self.connection)[1], FakeModel([json.dumps(dict(findings=[role]))]))
        second = FakeModel([json.dumps(dict(contexts=[], findings=[]))])
        self.extract(store, link_batches(self.connection, store, run_id)[0], second)
        self.assertIn("E1:F004", second.calls[0][1])
        self.assertIn("E1:F002", second.calls[0][1])
        self.assertNotIn("E1:F003", second.calls[0][1])

    def test_extraction_of_no_batches_makes_no_calls_at_any_worker_count(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        self.assertEqual(extract_batches(self.connection, store, [], call=FakeModel([]), model="fake-model",
                                         expected_fingerprints=self.expected, workers=3), [])

    def test_extraction_fails_closed_and_records_why(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        intro, cost = build_batches(self.connection)[1:3]
        fake = FakeModel([self.proposal(), "not json at all", RuntimeError("HTTP 429 rate limited"),
                          json.dumps(dict(contexts=[], findings=[]))])
        # The introduction batch was never shown the pricing passage: the context is kept, the finding is rejected.
        outcome, = self.extract(store, intro, fake)
        self.assertEqual((outcome.status, outcome.context_ids, outcome.finding_ids), ("findings", ["E1:C01"], []))
        self.assertEqual(len(outcome.rejected), 1)
        self.assertIn("E1:B0011", outcome.rejected[0])
        self.assertIsNone(outcome.error)
        for message in ("JSON", "429"):  # the reply or the call itself failed: nothing to keep
            with self.subTest(message=message):
                outcome, = self.extract(store, cost, fake)
                self.assertEqual((outcome.status, outcome.finding_ids, outcome.context_ids), ("failed", [], []))
                self.assertIn(message, outcome.error)
        self.assertEqual(self.extract(store, cost, fake)[0].status, "no_findings")
        self.assertEqual(self.extract(store, cost, fake)[0].status, "no_findings")  # reused, not re-asked
        self.assertEqual(len(fake.calls), 4)
        self.assertEqual(load_records(store, Finding), [])
        self.assertEqual([c.id for c in load_records(store, CompanyContext)], ["E1:C01"])
        run_record, = load_runs(store)
        self.assertEqual({o.batch_id: o.status for o in run_record.batches},
                         {"E1:S01": "findings", "E1:S02": "no_findings"})
        self.assertTrue(any(o.response == "not json at all" for o in run_record.batches) is False)  # latest outcome kept
        self.assertEqual(check_findings(store, self.connection), (1, 0))

    def test_rejected_records_are_listed_and_the_rest_are_kept(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        cost = build_batches(self.connection)[2]
        good = json.loads(self.proposal())["findings"][0]
        reply = dict(
            contexts=[dict(key="acme", company_name="Acme", employment="current", supporting=["E1:P006"]),
                      dict(key="ghost", company_name="Ghost", employment="current", supporting=["E1:P999"])],
            findings=[good, {**good, "supporting": ["E1:P010"]}, {**good, "context_ref": "ghost"},
                      {**good, "kind": "vendor_relationship", "relationship": None}])
        outcome, = self.extract(store, cost, FakeModel([json.dumps(reply)]))
        self.assertEqual((outcome.status, outcome.context_ids, outcome.finding_ids), ("findings", ["E1:C01"], ["E1:F001"]))
        self.assertEqual(len(outcome.rejected), 4)
        reasons = "\n".join(outcome.rejected)
        for expected in ("ghost", "E1:P999", "E1:B0010", "not supplied", "rejected context", "relationship"):
            self.assertIn(expected, reasons)
        self.assertEqual([f.id for f in load_records(store, Finding)], ["E1:F001"])
        self.assertEqual(check_findings(store, self.connection), (1, 1))

    def test_overlap_keeps_support_and_wording_when_qualifications_are_nonblank(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        reply = json.loads(self.proposal())
        proposed = reply["findings"][0]
        proposed.update(qualifying=["E1:P011", "E1:P009"],
                        qualifications=[" ", "The expert gives an approximate amount, not an exact invoice."])
        outcome, = self.extract(store, build_batches(self.connection)[2], FakeModel([json.dumps(reply)]))
        self.assertEqual((outcome.status, outcome.rejected), ("findings", []))
        finding, = load_records(store, Finding)
        self.assertEqual((finding.supporting, finding.qualifying, finding.context),
                         (["E1:B0011"], ["E1:B0009"], ["E1:B0006"]))
        self.assertEqual((finding.statement, finding.qualifications),
                         (proposed["statement"], proposed["qualifications"]))
        self.assertEqual(finding.review_state, "proposed")
        self.assertEqual(load_runs(store)[0].batches[0].response, json.dumps(reply))
        self.assertEqual(check_findings(store, self.connection), (1, 1))

    def test_overlap_normalization_does_not_bypass_citation_validation(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        batch = build_batches(self.connection)[2]
        supplied = set(batch.passage_ids + batch.context_ids)
        good = json.loads(self.proposal())
        for supporting, supplied_ids, message in (
                (["E1:P999"], supplied, "unknown citation"),
                (["E1:P011"], supplied - {"E1:B0011"}, "not supplied"),
                (["E1:P010"], supplied | {"E1:B0010"}, "text passage"),
                (["E1:P011", "E1:P011"], supplied, "one role")):
            with self.subTest(supporting=supporting, message=message):
                finding = {**good["findings"][0], "supporting": supporting, "qualifying": supporting[:1],
                           "qualifications": ["The expert estimates this amount."]}
                _, findings, rejected, _ = realize(
                    self.connection, store, batch, Proposal(contexts=good["contexts"], findings=[finding]),
                    SourceVersion(**self.version()), supplied_ids, self.expected)
                self.assertEqual(findings, [])
                self.assertIn(message, "\n".join(rejected))

    def test_current_context_warnings_identify_live_contexts_without_merging(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        contexts = [self.context(), self.context(id="E1:C02", company_name="Globex", employment="historical"),
                    self.context(id="E1:C03", company_name="Unclear Co", employment="unclear"),
                    self.context(id="E1:C04", company_name="Rejected Co", review_state="rejected")]
        save_run(store, self.run_record(), contexts, [self.finding()])
        self.assertEqual(extract.current_context_warnings(store), {})
        second = self.context(id="E1:C05", company_name="Second Job", review_state="reviewed",
                              note="The expert explicitly reports a second current job.")
        save_run(store, self.run_record(id="second-job"), [second], [])
        before = list(store.iterdump())
        warnings = extract.current_context_warnings(store)
        self.assertEqual(list(warnings), ["E1"])
        for text in ("E1:C01", "Acme", "E1:C05", "Second Job", "review"):
            self.assertIn(text, warnings["E1"])
        for text in ("E1:C02", "Globex", "E1:C03", "Unclear Co", "E1:C04", "Rejected Co"):
            self.assertNotIn(text, warnings["E1"])
        self.assertEqual(check_findings(store, self.connection), (5, 1))
        self.assertEqual(list(store.iterdump()), before)

    def test_lanes_match_sequential_results_and_keep_contexts_within_an_interview(self):
        directory = self.directory / "two"
        directory.mkdir()
        _, manifest = write_fixture(directory, experts=(1, 2))
        database = directory / "transcripts.sqlite"
        corpus, _ = run(manifest, database)
        expected = {document.id: document.extraction_sha256 for document in corpus.documents}
        connection = open_database(database, readonly=True)
        self.addCleanup(connection.close)
        batches = build_batches(connection)
        self.assertEqual([b.id for b in batches], ["E1:S00", "E1:S01", "E1:S02", "E2:S00", "E2:S01", "E2:S02"])
        summaries = {}
        for workers in (1, 2):
            store = open_findings(directory / f"findings-{workers}.sqlite")
            self.addCleanup(store.close)
            fake = LaneModel()
            outcomes = extract_batches(connection, store, batches, call=fake, model="fake-model",
                                       expected_fingerprints=expected, workers=workers)
            summaries[workers] = [(o.batch_id, o.status, o.finding_ids, o.context_ids) for o in outcomes]
            self.assertEqual(len(fake.calls), 6)
            cost_prompts = [user for user in fake.calls if "[E1:P011]" in user or "[E2:P011]" in user]
            self.assertTrue(all(re.search(r"- E\d:C01:", user) for user in cost_prompts),
                            "the cost batch must see the context its interview's introduction batch stored")
            self.assertEqual(all(fake.on_main_thread), workers == 1)
            self.assertEqual(check_findings(store, connection), (2, 2))
        self.assertEqual(summaries[1], summaries[2])
        self.assertEqual(summaries[1], [
            ("E1:S00", "no_findings", [], []), ("E1:S01", "findings", [], ["E1:C01"]),
            ("E1:S02", "findings", ["E1:F001"], []), ("E2:S00", "no_findings", [], []),
            ("E2:S01", "findings", [], ["E2:C01"]), ("E2:S02", "findings", ["E2:F001"], [])])

    def test_header_only_batches_get_no_model_call_because_every_batch_repeats_them(self):
        batches = build_batches(self.connection)
        self.assertEqual([b.id for b in model_batches(batches)], ["E1:S01", "E1:S02"])
        self.assertTrue(all("E1:B0001" in b.context_ids for b in model_batches(batches)))

    def test_a_later_section_adds_a_name_to_an_existing_company_only_with_a_passage_that_uses_it(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        header = self.context(company_name="Expert company", supporting=["E1:B0001"])  # the anonymized header
        save_run(store, self.run_record(), [header], [])
        cost = build_batches(self.connection)[2]
        reply = dict(contexts=[], findings=[], aliases=[
            dict(context_ref="E1:C01", alias="Acme", supporting=["E1:P006"]),  # "I lead IT at Acme, ..."
            dict(context_ref="E1:C01", alias="Initech", supporting=["E1:P006"])])  # not in the passage
        fake = FakeModel([json.dumps(reply)])
        outcome, = self.extract(store, cost, fake)
        self.assertEqual(outcome.alias_ids, ["E1:A01"])
        self.assertEqual(len(outcome.rejected), 1)
        self.assertIn("Initech", outcome.rejected[0])
        context, = load_records(store, CompanyContext)
        self.assertEqual((context.company_name, context.aliases), ("Expert company", ["Acme"]))
        self.assertEqual(check_findings(store, self.connection), (1, 0))
        again, = self.extract(store, cost, fake)  # its own alias is hidden from its input, so nothing changed
        self.assertEqual((len(fake.calls), again.alias_ids), (1, ["E1:A01"]))

    def test_review_records_decisions_and_an_edit_never_inherits_reviewed_status(self):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        save_run(store, self.run_record(), [self.context()], [self.finding()])
        reviewed = review(store, self.connection, "E1:F001", "reviewed", note="Checked against E1:P011.")
        self.assertEqual((reviewed.review_state, load_records(store, Finding)[0].review_state), ("reviewed", "reviewed"))
        edited = review(store, self.connection, "E1:F001", "reviewed", changes={"qualifying": ["E1:B0013"]},
                        note="The later question qualifies it.")
        self.assertEqual((edited.review_state, edited.qualifying), ("edited", ["E1:B0013"]))
        context = review(store, self.connection, "E1:C01", "reviewed", changes={"aliases": ["ACME Corp"]})
        self.assertEqual((context.review_state, context.aliases), ("edited", ["ACME Corp"]))
        with self.assertRaisesRegex(ValueError, "E2:B0004"):  # another interview's passage: refused, record unchanged
            review(store, self.connection, "E1:F001", "edited", changes={"supporting": ["E2:B0004"]})
        self.assertEqual(load_records(store, Finding)[0].supporting, ["E1:B0011"])
        for record_id, state, changes in (("E1:F999", "reviewed", None), ("E1:F001", "proposed", None),
                                          ("E1:F001", "reviewed", {"origin": "analyst"})):
            with self.subTest(record_id=record_id, state=state), self.assertRaises(ValueError):
                review(store, self.connection, record_id, state, changes=changes)
        rejected = review(store, self.connection, "E1:F001", "rejected", note="Duplicate of another finding.")
        self.assertEqual(rejected.review_state, "rejected")
        log = [tuple(row) for row in store.execute("SELECT record_id, state, note FROM reviews ORDER BY id")]
        self.assertEqual(log, [("E1:F001", "reviewed", "Checked against E1:P011."),
                               ("E1:F001", "edited", "The later question qualifies it."),
                               ("E1:C01", "edited", None), ("E1:F001", "rejected", "Duplicate of another finding.")])
        self.assertEqual(check_findings(store, self.connection), (1, 1))

    def test_env_loader_only_fills_missing_names(self):
        env_file = self.directory / "test.env"
        env_file.write_text('# comment\n\nXAI_API_KEY="from-file"\nOTHER=value \nexport SPACED = quoted\n', encoding="utf-8")
        with mock.patch.dict(os.environ, {"OTHER": "already-set"}, clear=False):
            os.environ.pop("XAI_API_KEY", None)
            self.assertEqual(load_env(env_file), ["XAI_API_KEY", "SPACED"])
            self.assertEqual((os.environ["XAI_API_KEY"], os.environ["OTHER"], os.environ["SPACED"]),
                             ("from-file", "already-set", "quoted"))
        self.assertEqual(load_env(self.directory / "missing.env"), [])

    def test_cli_extract_fails_without_credentials_and_stores_nothing(self):
        findings = self.directory / "findings.sqlite"
        environment = {key: value for key, value in os.environ.items() if key != "XAI_API_KEY"}
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("extract.py")),
                                 "--manifest", str(self.manifest), "--database", str(self.database),
                                 "--findings", str(findings), "--env", str(self.directory / "missing.env"),
                                 "--extract", "E1:S02"], capture_output=True, text=True, check=False, env=environment)
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("XAI_API_KEY", result.stderr)
        self.assertFalse(findings.exists())

    def test_cli_preflights_then_reports_batches_and_stale_findings(self):
        findings = self.directory / "findings.sqlite"

        def invoke(*arguments):
            return subprocess.run([sys.executable, str(Path(__file__).with_name("extract.py")),
                                   "--manifest", str(self.manifest), "--database", str(self.database),
                                   "--findings", str(findings), *arguments], capture_output=True, text=True, check=False)

        result = invoke("--batches")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("E1:S02", result.stdout)
        self.assertIn("3 batches cover 7 text passages and 4 chunks", result.stdout)
        result = invoke("--render", "E1:S02")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("[E1:P011] Expert 1 00:01:10: Roughly forty dollars per agent per month.", result.stdout)
        result = invoke("--check")
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("no findings database", result.stderr)
        store = open_findings(findings)
        save_run(store, self.run_record(), [self.context()], [self.finding()])
        store.close()
        result = invoke("--check")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("1 finding", result.stdout)
        store = open_findings(findings)
        save_run(store, self.run_record(id="second-context"),
                 [self.context(id="E1:C02", company_name="Acme Other Name")], [])
        store.close()
        before = findings.read_bytes()
        result = invoke("--check")
        self.assertEqual(result.returncode, 0, result.stderr)
        for text in ("WARNING E1:", "E1:C01", "Acme", "E1:C02", "Acme Other Name", "review"):
            self.assertIn(text, result.stdout)
        self.assertEqual(findings.read_bytes(), before)
        result = invoke("--findings", str(self.database), "--check")
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("separate file", result.stderr)
        connection = open_database(self.database)
        with connection:
            connection.execute("UPDATE passages SET text = 'Tampered' WHERE id = 'E1:B0011'")
        connection.close()
        for arguments in (("--batches",), ("--render", "E1:S02"), ("--check",)):
            with self.subTest(arguments=arguments):
                result = invoke(*arguments)
                self.assertEqual(result.returncode, 1, result.stdout)
                self.assertEqual(result.stdout, "")
                self.assertIn("Extraction failed:", result.stderr)

    def test_cli_evaluate_fails_closed_when_required_cases_are_absent(self):
        findings = self.directory / "empty-findings.sqlite"
        open_findings(findings).close()
        arguments = [str(Path(extract.__file__)), "--manifest", str(self.manifest), "--database",
                     str(self.database), "--findings", str(findings), "--evaluate"]
        version = {"E1": (self.document.source_sha256, self.document.extraction_sha256)}
        with mock.patch.dict(extract.GOLD_SOURCE_VERSIONS, version, clear=True):
            with mock.patch.object(sys, "argv", arguments):
                self.assertEqual(extract.main(), 1)

    def test_evaluation_refuses_transcripts_outside_the_reviewed_gold_version(self):
        findings = self.directory / "empty-findings.sqlite"
        store = open_findings(findings)
        self.addCleanup(store.close)
        result, = evaluate(store, self.connection)
        self.assertEqual(result[0:2], ("gold source versions", "fail"))


class AttributionTests(unittest.TestCase):
    """Which company a passage is about: the anchor rule, the batch input that shows it, and the check that uses it."""

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        corpus, _ = run(write_history_fixture(self.directory), self.directory / "transcripts.sqlite")
        self.document = corpus.documents[0]
        self.expected = {"E1": self.document.extraction_sha256}
        self.connection = open_database(self.directory / "transcripts.sqlite", readonly=True)
        self.addCleanup(self.connection.close)
        rows = self.connection.execute("SELECT id, text FROM passages WHERE kind = 'text'")
        self.ids = {text: passage_id for passage_id, text in rows}
        self.version = dict(source_sha256=self.document.source_sha256,
                            extraction_sha256=self.document.extraction_sha256)
        both = self.ids["I lead IT at Acme, previously at Globex."]
        self.contexts = [
            CompanyContext(id="E1:C01", document_id="E1", company_name="Acme", employment="current",
                           supporting=[both], origin="model", **self.version),
            CompanyContext(id="E1:C02", document_id="E1", company_name="Globex", employment="historical",
                           note="The expert's former employer.", supporting=[both], origin="model", **self.version)]

    def finding(self, context_id, text, **overrides):
        base = dict(id="E1:F001", document_id="E1", context_id=context_id, kind="pricing", vendor="ServiceNow",
                    statement="Total cost was roughly double the alternatives.", evidence_type="expert_estimate",
                    supporting=[self.ids[text]], origin="model", **self.version)
        return Finding(**{**base, **overrides})

    def store(self, findings=()):
        store = open_findings(self.directory / "findings.sqlite")
        self.addCleanup(store.close)
        save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                      source_versions={"E1": SourceVersion(**self.version)}),
                 self.contexts, list(findings))
        return store

    def test_anchor_follows_the_latest_passage_naming_exactly_one_company(self):
        anchors = company_anchors(self.connection, self.contexts, "E1")
        globex = self.ids["What did Globex run?"]
        self.assertEqual({text: anchors[i][:2] for text, i in self.ids.items()}, {
            "Expert 1": ("E1:C01", None),  # the interview starts at the current employer
            "Tell me about your role.": ("E1:C01", None),
            "I lead IT at Acme, previously at Globex.": ("E1:C01", None),  # naming both moves nothing
            "What did Globex run?": ("E1:C02", globex),  # an interviewer question can set what is discussed
            "We ran ServiceNow there for years.": ("E1:C02", globex),
            "How did total cost compare?": ("E1:C02", globex),
            "Roughly double the alternatives.": ("E1:C02", globex)})
        self.assertEqual(company_anchors(self.connection, [], "E1"), {})

    def test_batch_input_quotes_company_profiles_and_the_conversation_anchor_and_makes_them_citable(self):
        cost = build_batches(self.connection)[3]
        rendered = render_batch(self.connection, cost, expected_fingerprints=self.expected)
        text, extra = user_message(self.connection, self.contexts, cost, rendered, expected_fingerprints=self.expected)
        self.assertEqual(extra, [self.ids["What did Globex run?"], self.ids["We ran ServiceNow there for years."]])
        self.assertIn("- E1:C02: Globex (historical); note: The expert's former employer.", text)
        self.assertIn("Most recent explicit company reference before this section: E1:C02 Globex at E1:P009", text)
        self.assertIn("[E1:P011] Expert 1 00:00:05: We ran ServiceNow there for years.", text)
        self.assertTrue(text.endswith(rendered))
        # The model can now cite the passage that establishes Globex; the old input made that citation unsupplied.
        store = self.store()  # the contexts earlier batches stored
        reply = dict(contexts=[], findings=[
            dict(context_ref="E1:C02", kind="pricing", vendor="ServiceNow", evidence_type="expert_estimate",
                 statement="Total cost was roughly double the alternatives.", supporting=["E1:P016"],
                 context=["E1:P009", "E1:P016"])])  # a context copy of a supporting passage is dropped, not fatal
        outcome, = extract_batches(self.connection, store, [cost], call=FakeModel([json.dumps(reply)]),
                                   model="fake-model", expected_fingerprints=self.expected)
        self.assertEqual((outcome.status, outcome.rejected), ("findings", []))
        finding, = load_records(store, Finding)
        self.assertEqual((finding.supporting, finding.context),
                         ([self.ids["Roughly double the alternatives."]], [self.ids["What did Globex run?"]]))
        self.assertTrue(set(extra) <= set(outcome.input_passage_ids))

    def test_input_fingerprint_covers_the_company_profiles(self):
        cost = build_batches(self.connection)[3]
        prints = []
        for note in ("first", "second"):
            store = open_findings(self.directory / f"findings-{note}.sqlite")
            self.addCleanup(store.close)
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                          source_versions={"E1": SourceVersion(**self.version)}),
                     [self.contexts[0], self.contexts[1].model_copy(update={"note": note})], [])
            empty = FakeModel(['{"contexts": [], "findings": []}'])
            outcome, = extract_batches(self.connection, store, [cost], call=empty,
                                       model="fake-model", expected_fingerprints=self.expected)
            prints.append(outcome.input_fingerprint)
        self.assertNotEqual(*prints)

    def test_attribution_check_flags_a_finding_filed_against_the_conversation_anchor(self):
        store = self.store([
            self.finding("E1:C01", "Roughly double the alternatives."),  # filed to Acme while Globex is discussed
            self.finding("E1:C02", "Roughly double the alternatives.", id="E1:F002"),
            self.finding("E1:C01", "I lead IT at Acme, previously at Globex.", id="E1:F003"),  # names its company
            self.finding(None, "Roughly double the alternatives.", id="E1:F004")])  # left unestablished: no dispute
        flags = attribution_flags(store, self.connection)
        self.assertEqual(list(flags), ["E1:F001"])
        self.assertIn("Globex", flags["E1:F001"])
        self.assertIn("E1:P009", flags["E1:F001"])

    def test_a_finding_may_leave_its_company_unestablished_until_a_person_sets_it(self):
        store = self.store([self.finding(None, "Roughly double the alternatives.")])
        self.assertEqual(check_findings(store, self.connection), (2, 1))
        with self.assertRaisesRegex(ValueError, "company"):
            review(store, self.connection, "E1:F001", "reviewed")
        edited = review(store, self.connection, "E1:F001", "edited", changes={"context_id": "E1:C02"},
                        note="The cost section continues the Globex account (E1:P009).")
        self.assertEqual((edited.context_id, edited.review_state), ("E1:C02", "edited"))
        self.assertEqual(review(store, self.connection, "E1:F001", "reviewed").review_state, "reviewed")
        for context_id in ("E1:C99", "E2:C01"):  # missing, or another interview's: refused, record unchanged
            with self.subTest(context_id=context_id), self.assertRaisesRegex(ValueError, "company context"):
                review(store, self.connection, "E1:F001", "edited", changes={"context_id": context_id})
        self.assertEqual(load_records(store, Finding)[0].context_id, "E1:C02")
        # A record whose context disappears from the store fails the check instead of passing silently.
        store.execute("PRAGMA foreign_keys = OFF")  # a no-op inside a transaction, so before it
        with store:
            store.execute("DELETE FROM company_contexts WHERE id = 'E1:C02'")
        with self.assertRaisesRegex(ValueError, "company context E1:C02"):
            check_findings(store, self.connection)

    def test_interviewer_text_moves_to_context_and_cannot_support_alone(self):
        store = self.store()
        cost = build_batches(self.connection)[3]
        base = dict(context_ref="E1:C02", kind="pricing", evidence_type="expert_estimate",
                    statement="Total cost was roughly double the alternatives.")
        reply = dict(contexts=[], findings=[
            {**base, "supporting": ["E1:P014", "E1:P016"]},  # the question plus the expert's answer
            {**base, "statement": "Cost was compared.", "supporting": ["E1:P014"]}])  # the question alone
        outcome, = extract_batches(self.connection, store, [cost], call=FakeModel([json.dumps(reply)]),
                                   model="fake-model", expected_fingerprints=self.expected)
        finding, = load_records(store, Finding)
        self.assertEqual((finding.supporting, finding.context),
                         ([self.ids["Roughly double the alternatives."]], [self.ids["How did total cost compare?"]]))
        self.assertEqual(len(outcome.rejected), 1)
        self.assertIn("interviewer", outcome.rejected[0])

    def test_overlap_without_nonblank_qualifications_is_rejected(self):
        store = self.store()
        cost = build_batches(self.connection)[3]
        reply = dict(contexts=[], findings=[dict(context_ref="E1:C02", kind="pricing", evidence_type="expert_estimate",
                                                 statement="Total cost was roughly double.", supporting=["E1:P016"],
                                                 qualifying=["E1:P016"])])
        for index, qualifications in enumerate((None, [], [""], [" ", "\t\n"])):
            with self.subTest(qualifications=qualifications):
                if qualifications is not None:
                    reply["findings"][0]["qualifications"] = qualifications
                outcome, = extract_batches(self.connection, store, [cost], call=FakeModel([json.dumps(reply)]),
                                           model="fake-model", prompt_version=f"empty-qualification-{index}",
                                           expected_fingerprints=self.expected)
                self.assertEqual((outcome.status, outcome.finding_ids), ("no_findings", []))
                self.assertIn("one role", outcome.rejected[0])


class QuantityNormalizationTests(unittest.TestCase):
    """Prices are compared inside one answer: currency normalized to dollars, unit and period carried only there."""

    def test_only_a_price_rate_borrows_currency_unit_and_period_from_its_answer(self):
        answer = ("We're on a tier that comes out to roughly forty dollars per agent per month. ServiceNow's quote "
                  "for a comparable scope was north of one-fifty. Ivanti was around fifty-five.")
        carried = normalize_quantity(Quantity(raw="north of one-fifty", low=150.0), answer, kind="pricing")
        self.assertEqual((carried.currency, carried.unit, carried.period, carried.carried),
                         ("dollars", "per agent", "month", ["currency", "unit", "period"]))
        stated = normalize_quantity(Quantity(raw="about $100 per user/month", value=100.0, currency="USD",
                                             unit="per user", period="month"), "about $100 per user/month",
                                    kind="pricing")
        self.assertEqual((stated.currency, stated.carried), ("dollars", []))  # stated, only relabeled
        dollar_sign = normalize_quantity(Quantity(raw="roughly $50-55", low=50.0, high=55.0),
                                         "ServiceNow at about $100 per user/month; Zendesk at roughly $50-55.",
                                         kind="pricing")
        self.assertEqual((dollar_sign.currency, dollar_sign.unit, dollar_sign.period, dollar_sign.carried),
                         ("dollars", "per user", "month", ["unit", "period"]))
        blended = normalize_quantity(Quantity(raw="around seventy", value=70.0),
                                     "BMC at roughly sixty dollars a user per month; Ivanti was around seventy.",
                                     kind="pricing")
        self.assertEqual((blended.unit, blended.period), ("per user", "month"))
        mixed = normalize_quantity(Quantity(raw="around seventy", value=70.0),
                                   "Sixty dollars per user per month, forty dollars per agent per month; "
                                   "around seventy.",
                                   kind="pricing")
        self.assertEqual((mixed.currency, mixed.unit, mixed.carried), ("dollars", None, ["currency"]))  # units differ
        ratio = Quantity(raw="almost double")  # a comparison, not a price: nothing to borrow
        self.assertEqual(normalize_quantity(ratio, "ServiceNow at about $100 per user/month.", kind="pricing"), ratio)

    def test_a_figure_that_is_not_a_price_borrows_nothing_and_takes_its_unit_from_its_own_words(self):
        company = "about a $40B company with roughly 2,000 sites"
        sites = normalize_quantity(Quantity(raw="roughly 2,000 sites", value=2000.0), company, kind="company_scale")
        self.assertEqual((sites.currency, sites.carried), (None, []))  # a count never takes the answer's dollars
        duration = normalize_quantity(Quantity(raw="about seven months", value=7.0, basis="kickoff to go-live"),
                                      "It took about seven months from kickoff to go-live.", kind="implementation")
        self.assertEqual((duration.value, duration.unit, duration.raw, duration.carried),
                         (7.0, "months", "about seven months", []))
        expected = normalize_quantity(Quantity(raw="a 6-7 month implementation", low=6.0, high=7.0),
                                      "we expected a 6-7 month implementation", kind="implementation")
        self.assertEqual(expected.unit, "months")
        share = normalize_quantity(Quantity(raw="around 30-40% of their time", low=30.0, high=40.0),
                                   "3-5 employees, each allocated around 30-40% of their time", kind="implementation")
        self.assertEqual((share.unit, share.currency), ("percent", None))
        renewal = normalize_quantity(Quantity(raw="6 to 8 percent", low=6.0, high=8.0),
                                     "Forty dollars per agent per month, and renewals rose 6 to 8 percent.",
                                     kind="pricing")
        self.assertEqual((renewal.unit, renewal.currency, renewal.carried), ("percent", None, []))
        stated = Quantity(raw="about 3-5 employees", low=3.0, high=5.0, unit="employees")
        self.assertEqual(normalize_quantity(stated, "about 3-5 employees", kind="implementation"), stated)

    def test_a_rating_answered_on_another_scale_is_labeled_with_the_question(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            document = Document()
            document.add_paragraph("Expert 1")
            document.add_paragraph("Renewal", style="Heading 1")
            document.add_paragraph("AI Interviewer  00:00:01")
            document.add_paragraph("On a scale of 1 to 7, how likely are you to renew?")
            document.add_paragraph("Expert 1  00:00:05")
            document.add_paragraph("I would say 9 out of 10.")
            path = directory / "rating.docx"
            document.save(path)
            manifest = directory / "manifest.json"
            manifest.write_text(json.dumps([Source(expert=1, source=path.name, sha256=hashlib.sha256(
                path.read_bytes()).hexdigest()).model_dump()]), encoding="utf-8")
            corpus, _ = run(manifest, directory / "transcripts.sqlite")
            expected = {"E1": corpus.documents[0].extraction_sha256}
            connection = open_database(directory / "transcripts.sqlite", readonly=True)
            store = open_findings(directory / "findings.sqlite")
            batch = build_batches(connection)[1]
            reply = dict(contexts=[dict(key="acme", company_name="Acme", employment="current",
                                        supporting=["E1:P004"])],
                         findings=[dict(context_ref="acme", kind="selection_criteria", evidence_type="opinion",
                                        statement="The expert rates renewal likelihood 9 out of 10.",
                                        supporting=["E1:P006"])])
            outcome, = extract_batches(connection, store, [batch], call=FakeModel([json.dumps(reply)]),
                                       model="fake-model", expected_fingerprints=expected)
            finding, = load_records(store, Finding)
            store.close()
            connection.close()
        self.assertEqual(outcome.rejected, [])
        self.assertIn("E1:B0004", finding.context)
        self.assertTrue(any("1 to 7" in q and "9 out of 10" in q for q in finding.qualifications))


@unittest.skipUnless(MANIFEST.is_file(), "Private transcript manifest is not configured; synthetic tests still run.")
class ProvidedTranscriptExtractionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.corpus, _ = run(MANIFEST, DEFAULT_DATABASE, check=True)  # fails, not skips, on missing or stale files
        cls.expected = {document.id: document.extraction_sha256 for document in cls.corpus.documents}
        cls.connection = open_database(DEFAULT_DATABASE, readonly=True)
        cls.batches = {batch.id: batch for batch in build_batches(cls.connection)}

    @classmethod
    def tearDownClass(cls):
        cls.connection.close()

    def test_schedule_covers_all_three_interviews_section_by_section(self):
        self.assertEqual(len(self.batches), 30)
        self.assertEqual(sum(len(b.passage_ids) for b in self.batches.values()), 315)
        self.assertEqual(sum(len(b.chunk_ids) for b in self.batches.values()), 160)
        self.assertEqual(self.batches["E1:S03"].heading, "Vendor Selection & Decision Criteria")
        self.assertEqual(len(self.batches["E1:S03"].passage_ids), 38)
        self.assertEqual(self.batches["E1:S02"].heading, "Current ServiceNow Environment")
        text = render_batch(self.connection, self.batches["E1:S02"], expected_fingerprints=self.expected)
        before, after = text.split("Section passages (primary evidence):")
        self.assertIn("Section heading (navigation only, not evidence): Current ServiceNow Environment", before)
        self.assertIn("[E1:P016] Expert 1 ", before)  # the current-employer opening travels as labeled context
        self.assertIn("[E1:P021] Expert 1 ", after)  # the historical Thermo Fisher account is this section's evidence
        self.assertIn("Thermo Fisher", after)

    def test_solara_pricing_finding_validates_stores_and_checks(self):
        e3 = next(document for document in self.corpus.documents if document.id == "E3")
        version = dict(source_sha256=e3.source_sha256, extraction_sha256=e3.extraction_sha256)
        context = CompanyContext(id="E3:C01", document_id="E3", company_name="Solara Renewables", employment="current",
                                 supporting=["E3:B0017"], origin="analyst", **version)
        finding = Finding(
            id="E3:F001", document_id="E3", context_id="E3:C01", kind="pricing", vendor="Freshservice",
            statement="The expert reports roughly forty dollars per agent per month for Solara's Freshservice tier.",
            evidence_type="expert_estimate",
            quantity=Quantity(raw="roughly forty dollars per agent per month", value=40.0, approximate=True,
                              unit="per agent", period="month", currency="dollars", basis="current tier"),
            supporting=["E3:B0098"], qualifying=["E3:B0096"], context=["E3:B0017"], origin="analyst", **version)
        batch = self.batches["E3:S06"]
        self.assertEqual(batch.heading, "Cost & Total Cost of Ownership")
        rows = validate_record(self.connection, finding, set(batch.passage_ids) | set(batch.context_ids),
                               expected_fingerprints=self.expected)
        self.assertEqual([r["citation_id"] for r in rows], ["E3:P017", "E3:P096", "E3:P098"])
        self.assertIn("roughly forty dollars per agent per month", rows[2]["text"])
        with tempfile.TemporaryDirectory() as temporary:
            store = open_findings(Path(temporary) / "findings.sqlite")
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                          source_versions={"E3": SourceVersion(**version)}), [context], [finding])
            self.assertEqual(check_findings(store, self.connection), (1, 1))
            self.assertEqual(load_records(store, Finding)[0].quantity.unit, "per agent")
            store.close()

    def test_evaluation_distinguishes_correct_records_from_known_failures(self):
        documents = {d.id: d for d in self.corpus.documents}
        version = {d: dict(source_sha256=documents[d].source_sha256, extraction_sha256=documents[d].extraction_sha256)
                   for d in ("E1", "E2")}
        physical = {row[1]: row[0] for row in self.connection.execute(
            "SELECT id, citation_id FROM passages WHERE citation_id IS NOT NULL")}

        def finding(number, context_id, supporting, **extra):
            document = context_id.split(":")[0] if context_id else "E1"
            base = dict(id=f"{document}:F{number:03d}", document_id=document, context_id=context_id,
                        kind="implementation", statement="A statement.", evidence_type="firsthand_report",
                        supporting=[physical[c] for c in supporting], origin="analyst", **version[document])
            if "context" in extra:
                extra["context"] = [physical[c] for c in extra["context"]]
            return Finding(**{**base, **extra})

        contexts = [CompanyContext(id="E1:C01", document_id="E1", company_name="Pharma Company", employment="current",
                                   supporting=[physical["E1:P001"]], origin="analyst", **version["E1"]),
                    CompanyContext(id="E1:C02", document_id="E1", company_name="Thermo Fisher", employment="historical",
                                   supporting=[physical["E1:P021"]], origin="analyst", **version["E1"])]
        findings = [
            finding(1, "E1:C01", ["E1:P201"], kind="pricing"),  # cost section filed to the current employer
            finding(2, "E1:C02", ["E1:P176"], quantity=Quantity(raw="roughly 9-12 months", low=9, high=12,
                                                                  approximate=True, unit="months")),
            finding(3, "E1:C02", ["E1:P192"], quantity=Quantity(raw="a 6-7 month implementation", low=6, high=7,
                                                                unit="months")),
            finding(6, "E1:C02", ["E1:P176"]),  # the pain points from the same passage: not the duration
            finding(4, "E1:C02", ["E1:P237"], kind="pricing"),  # the bare yes, without its clarification
            finding(5, "E1:C01", ["E1:P282"], kind="vendor_relationship", vendor="ServiceNow",
                    relationship="previously_used"),
        ]
        with tempfile.TemporaryDirectory() as temporary:
            store = open_findings(Path(temporary) / "findings.sqlite")
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed", source_versions={
                d: SourceVersion(**v) for d, v in version.items()}), contexts, findings)
            results = {case: (status, detail) for case, status, detail in evaluate(store, self.connection)}
            store.close()
        self.assertEqual(results["G03 Thermo cost is implementation TCO"][0], "fail")
        self.assertIn("E1:F001", results["G03 Thermo cost is implementation TCO"][1])
        self.assertNotIn("E1:F002", results["G03 Thermo cost is implementation TCO"][1])
        self.assertEqual(results["G09 testing-access clarification qualifies the P237 yes"][0], "fail")
        self.assertIn("E1:F004", results["G09 testing-access clarification qualifies the P237 yes"][1])
        self.assertEqual(results["G12 Thermo expected and actual rollout durations preserve months"][0], "pass")
        self.assertEqual(results["G13 BMC timing tension remains linked and unreconciled"][0], "not_run")
        self.assertNotEqual(results["G05 module status stays installed/evaluated/considered/unpurchased"][0], "fail")
        store = open_findings(Path(tempfile.mkdtemp()) / "findings.sqlite")
        self.addCleanup(store.close)
        save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed", source_versions={
            d: SourceVersion(**v) for d, v in version.items()}), contexts, [findings[2], findings[3]])
        rollout = dict((case, (status, detail)) for case, status, detail in evaluate(store, self.connection))[
            "G12 Thermo expected and actual rollout durations preserve months"]
        self.assertEqual(rollout[0], "fail")  # P176 cited, but no finding states its 9-12 month duration
        self.assertIn("9-12", rollout[1])

        historical_cost = finding(
            7, "E1:C02", ["E1:P201"], kind="pricing", vendor="ServiceNow",
            statement="ServiceNow had higher implementation costs than competitors.",
            quantity=Quantity(raw="implementation cost", basis="implementation total cost of ownership"))
        cost_store = open_findings(Path(tempfile.mkdtemp()) / "missing-comparison.sqlite")
        self.addCleanup(cost_store.close)
        save_run(cost_store, ExtractionRun(id="cost-seed", model="none", prompt_version="analyst-seed",
                                           source_versions={d: SourceVersion(**v) for d, v in version.items()}),
                 contexts, [historical_cost])
        cost_case = dict((case, (status, detail)) for case, status, detail in evaluate(
            cost_store, self.connection))["G03 Thermo cost is implementation TCO"]
        self.assertEqual(cost_case[0], "fail")
        self.assertIn("roughly double", cost_case[1])

    def test_raw_evaluation_scores_unedited_model_output_only(self):
        e1 = next(document for document in self.corpus.documents if document.id == "E1")
        version = dict(source_sha256=e1.source_sha256, extraction_sha256=e1.extraction_sha256)
        physical = dict(self.connection.execute("SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        contexts = [CompanyContext(id="E1:C01", document_id="E1", company_name="Pharma Company", employment="current",
                                   supporting=[physical["E1:P001"]], origin="model", **version),
                    CompanyContext(id="E1:C02", document_id="E1", company_name="Thermo Fisher", employment="historical",
                                   supporting=[physical["E1:P021"]], origin="model", **version)]
        base = dict(document_id="E1", kind="pricing", vendor="ServiceNow", evidence_type="expert_estimate",
                    statement="Implementation TCO was roughly double.", supporting=[physical["E1:P201"]], **version)
        findings = [Finding(id="E1:F001", context_id="E1:C01", origin="model", **base),  # the model's miss
                    Finding(id="E1:F002", context_id="E1:C02", origin="model", review_state="edited", **base),
                    Finding(id="E1:F003", context_id="E1:C02", origin="analyst", **base)]
        with tempfile.TemporaryDirectory() as temporary:
            store = open_findings(Path(temporary) / "findings.sqlite")
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                          source_versions={"E1": SourceVersion(**version)}), contexts, findings)
            case = "G03 Thermo cost is implementation TCO"
            edited = {name: detail for name, _, detail in evaluate(store, self.connection)}[case]
            raw = {name: detail for name, _, detail in evaluate(store, self.connection, raw=True)}[case]
            store.close()
        self.assertIn("E1:F001", edited)
        self.assertIn("E1:F001", raw)
        self.assertIn("E1:F002", edited)
        self.assertNotIn("E1:F002", raw)  # an analyst edit is not model output
        self.assertNotIn("E1:F003", raw)

    def test_price_guard_rejects_currency_normalization(self):
        e3 = next(document for document in self.corpus.documents if document.id == "E3")
        version = dict(source_sha256=e3.source_sha256, extraction_sha256=e3.extraction_sha256)
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        context = CompanyContext(id="E3:C01", document_id="E3", company_name="Solara Renewables",
                                 employment="current", supporting=[physical["E3:P017"]], origin="analyst",
                                 **version)
        finding = Finding(
            id="E3:F001", document_id="E3", context_id="E3:C01", kind="pricing", vendor="ServiceNow",
            statement="The ServiceNow quote was north of one-fifty dollars per agent per month.",
            evidence_type="expert_estimate",
            quantity=Quantity(raw="north of one-fifty", low=150.0, unit="per agent", period="month",
                              currency="USD", basis="vendor quote"),
            supporting=[physical["E3:P098"]], context=[physical["E3:P017"]], origin="analyst", **version)
        with tempfile.TemporaryDirectory() as temporary:
            store = open_findings(Path(temporary) / "findings.sqlite")
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                          source_versions={"E3": SourceVersion(**version)}), [context], [finding])
            result = {name: (status, detail) for name, status, detail in evaluate(store, self.connection)}[
                "G14 license prices retain company, amount, unit, currency, and stated basis"]
            store.close()
        self.assertEqual(result[0], "fail")
        self.assertIn("E3:F001 must preserve the source currency wording 'dollars'", result[1])

    def test_cost_and_switching_findings_are_filed_to_thermo_fisher_or_flagged(self):
        """E1:P195-P249 (cost) and E1:P284-P311 (switching) continue the Thermo Fisher account."""
        store = open_findings(DEFAULT_FINDINGS, readonly=True)
        self.addCleanup(store.close)
        citation = {row[0]: row[1] for row in self.connection.execute(
            "SELECT id, citation_id FROM passages WHERE document_id = 'E1' AND citation_id IS NOT NULL")}
        names = {c.id: c.company_name for c in load_records(store, CompanyContext)}
        flags = attribution_flags(store, self.connection)
        in_scope = [f for f in load_records(store, Finding) if f.document_id == "E1" and f.review_state != "rejected"
                    and any(195 <= (n := int(citation[i].rsplit("P", 1)[1])) <= 249 or 284 <= n
                            <= 311 for i in f.supporting)]
        self.assertTrue(in_scope)
        unflagged_elsewhere = [f.id for f in in_scope if "Thermo" not in names.get(f.context_id or "", "")
                               and f.id not in flags]
        self.assertEqual(unflagged_elsewhere, [])

    def test_evaluation_rejects_interviewer_prompt_used_as_support(self):
        e1 = next(document for document in self.corpus.documents if document.id == "E1")
        version = dict(source_sha256=e1.source_sha256, extraction_sha256=e1.extraction_sha256)
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        context = CompanyContext(id="E1:C01", document_id="E1", company_name="Thermo Fisher",
                                 employment="historical", supporting=[physical["E1:P021"]], origin="analyst",
                                 **version)
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)

        def seeded(name, supporting, context_citations):
            finding = Finding(
                id="E1:F001", document_id="E1", context_id="E1:C01", kind="pricing", vendor="Freshservice",
                statement="The expert agreed that Freshservice wins on lower cost and faster deployment.",
                evidence_type="prompted_agreement", supporting=[physical[c] for c in supporting],
                context=[physical[c] for c in context_citations], origin="analyst", **version)
            store = open_findings(Path(temporary.name) / f"{name}.sqlite")
            save_run(store, ExtractionRun(id="seed", model="none", prompt_version="analyst-seed",
                                          source_versions={"E1": SourceVersion(**version)}), [context], [finding])
            self.addCleanup(store.close)
            return store

        invalid = seeded("prompt-in-support", ["E1:P145", "E1:P147"], ["E1:P021"])
        invalid_results = {name: (status, detail) for name, status, detail in evaluate(invalid, self.connection)}
        self.assertEqual(invalid_results["G07/G08 interviewer premises are not supporting evidence"][0], "fail")
        self.assertIn("E1:F001", invalid_results["G07/G08 interviewer premises are not supporting evidence"][1])

        valid = seeded("prompt-as-context", ["E1:P147"], ["E1:P021", "E1:P145"])
        valid_results = {name: status for name, status, _ in evaluate(valid, self.connection)}
        # not a failure; and not a pass either, because this seed covers one interview, not every section
        self.assertEqual(valid_results["G07/G08 interviewer premises are not supporting evidence"], "not_run")

    def test_ivanti_guard_requires_product_and_prior_deployment_details(self):
        documents = {document.id: document for document in self.corpus.documents}
        versions = {document_id: (document.source_sha256, document.extraction_sha256)
                    for document_id, document in documents.items()}
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        calloway = CompanyContext(
            id="E2:C01", document_id="E2", company_name="Calloway Biosciences", employment="current",
            supporting=[physical["E2:P015"]], origin="analyst",
            source_sha256=versions["E2"][0], extraction_sha256=versions["E2"][1])
        solara = CompanyContext(
            id="E3:C01", document_id="E3", company_name="Solara Renewables", employment="current",
            supporting=[physical["E3:P017"]], origin="analyst",
            source_sha256=versions["E3"][0], extraction_sha256=versions["E3"][1])
        previous_employer = CompanyContext(
            id="E3:C02", document_id="E3", company_name="Industrial Services Company",
            employment="historical", supporting=[physical["E3:P017"]], origin="analyst",
            source_sha256=versions["E3"][0], extraction_sha256=versions["E3"][1])
        rows = {
            "E2": {"id": "E2:F001", "document_id": "E2", "context_id": "E2:C01", "kind": "vendor_relationship",
                   "vendor": "Ivanti Neurons", "statement": "Calloway evaluated Ivanti Neurons.",
                   "relationship": "evaluated", "evidence_type": "firsthand_report",
                   "supporting": [physical["E2:P039"]],
                   "origin": "analyst", "source_sha256": versions["E2"][0], "extraction_sha256": versions["E2"][1]},
            "E3": {"id": "E3:F001", "document_id": "E3", "context_id": "E3:C02", "kind": "vendor_relationship",
                   "vendor": "Ivanti Service Manager", "statement": "A previous employer ran Ivanti Service Manager.",
                   "relationship": "previously_used", "evidence_type": "firsthand_report",
                   "supporting": [physical["E3:P017"]], "origin": "analyst",
                   "source_sha256": versions["E3"][0], "extraction_sha256": versions["E3"][1]},
        }
        run = ExtractionRun(id="seed", model="none", prompt_version="analyst-seed", source_versions={
            document_id: SourceVersion(source_sha256=version[0], extraction_sha256=version[1])
            for document_id, version in versions.items()})
        with tempfile.TemporaryDirectory() as temporary, mock.patch.dict(
                extract.GOLD_SOURCE_VERSIONS, versions, clear=True):
            store = open_findings(Path(temporary) / "findings.sqlite")
            self.addCleanup(store.close)
            records = [Finding(**row) for row in rows.values()]
            save_run(store, run, [calloway, solara, previous_employer], records)
            result = {name: status for name, status, _ in evaluate(store, self.connection)}[
                "G17 Ivanti products and experience are not conflated"]
            self.assertEqual(result, "fail")

            rows["E3"]["statement"] = (
                "At the previous employer, the organization ran Ivanti Service Manager, formerly Ivanti Heat, "
                "on-prem.")
            corrected = [Finding(**row) for row in rows.values()]
            replacement = open_findings(Path(temporary) / "corrected.sqlite")
            self.addCleanup(replacement.close)
            save_run(replacement, run.model_copy(update={"id": "corrected"}),
                     [calloway, solara, previous_employer], corrected)
            result = {name: status for name, status, _ in evaluate(replacement, self.connection)}[
                "G17 Ivanti products and experience are not conflated"]
            self.assertEqual(result, "pass")

    def test_future_servicenow_guard_requires_both_conditional_company_cases(self):
        documents = {document.id: document for document in self.corpus.documents}
        versions = {document_id: (document.source_sha256, document.extraction_sha256)
                    for document_id, document in documents.items()}
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        contexts = [
            CompanyContext(id="E2:C01", document_id="E2", company_name="Calloway Biosciences",
                           employment="current", supporting=[physical["E2:P015"]], origin="analyst",
                           source_sha256=versions["E2"][0], extraction_sha256=versions["E2"][1]),
            CompanyContext(id="E3:C01", document_id="E3", company_name="Solara Renewables",
                           employment="current", supporting=[physical["E3:P017"]], origin="analyst",
                           source_sha256=versions["E3"][0], extraction_sha256=versions["E3"][1]),
        ]

        def record(document_id, context_id, citation_id, statement):
            return Finding(
                id=f"{document_id}:F001", document_id=document_id, context_id=context_id,
                kind="vendor_relationship", vendor="ServiceNow", statement=statement,
                relationship="hypothetical", evidence_type="hypothetical",
                supporting=[physical[citation_id]], origin="analyst",
                source_sha256=versions[document_id][0], extraction_sha256=versions[document_id][1])

        calloway = record("E2", "E2:C01", "E2:P156",
                          "If Calloway switched, the expert would evaluate ServiceNow first.")
        solara = record("E3", "E3:C01", "E3:P140",
                        "If Solara tripled in size, the expert could consider ServiceNow.")
        run = ExtractionRun(id="seed", model="none", prompt_version="analyst-seed", source_versions={
            document_id: SourceVersion(source_sha256=version[0], extraction_sha256=version[1])
            for document_id, version in versions.items()})
        with tempfile.TemporaryDirectory() as temporary, mock.patch.dict(
                extract.GOLD_SOURCE_VERSIONS, versions, clear=True):
            for name, records, expected in (
                    ("complete", [calloway, solara], "pass"),
                    ("missing-solara", [calloway], "fail"),
                    ("asserted-move", [calloway, solara.model_copy(update={
                        "statement": "Solara is moving to ServiceNow."})], "fail")):
                with self.subTest(name=name):
                    store = open_findings(Path(temporary) / f"{name}.sqlite")
                    self.addCleanup(store.close)
                    save_run(store, run.model_copy(update={"id": name}), contexts, records)
                    result = {case: status for case, status, _ in evaluate(store, self.connection)}[
                        "G06 future ServiceNow moves stay conditional"]
                    self.assertEqual(result, expected)

    def test_testing_environment_answer_must_preserve_user_access_clarification(self):
        e1 = next(document for document in self.corpus.documents if document.id == "E1")
        versions = {document.id: (document.source_sha256, document.extraction_sha256)
                    for document in self.corpus.documents}
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        context = CompanyContext(
            id="E1:C01", document_id="E1", company_name="Thermo Fisher", employment="historical",
            supporting=[physical["E1:P021"]], origin="analyst",
            source_sha256=e1.source_sha256, extraction_sha256=e1.extraction_sha256)

        def record(statement, qualifications):
            return Finding(
                id="E1:F001", document_id="E1", context_id="E1:C01", kind="implementation",
                statement=statement, evidence_type="firsthand_report",
                supporting=[physical["E1:P237"], physical["E1:P245"]], qualifications=qualifications,
                context=[physical["E1:P021"]], origin="analyst",
                source_sha256=e1.source_sha256, extraction_sha256=e1.extraction_sha256)

        clarify = ["The clarification was about business users needing access to test workflows, not buying "
                   "entirely separate non-production environments."]
        correct = record("Testing users were underestimated.", clarify)
        incorrect = record("Thermo purchased separate non-production environments for testing.", clarify)
        run = ExtractionRun(id="seed", model="none", prompt_version="analyst-seed", source_versions={
            document_id: SourceVersion(source_sha256=version[0], extraction_sha256=version[1])
            for document_id, version in versions.items()})
        with tempfile.TemporaryDirectory() as temporary, mock.patch.dict(
                extract.GOLD_SOURCE_VERSIONS, versions, clear=True):
            for name, finding, expected in (("correct", correct, "pass"), ("purchase-claim", incorrect, "fail")):
                with self.subTest(name=name):
                    store = open_findings(Path(temporary) / f"{name}.sqlite")
                    self.addCleanup(store.close)
                    save_run(store, run.model_copy(update={"id": name}), [context], [finding])
                    result = {case: (status, detail) for case, status, detail in evaluate(store, self.connection)}[
                        "G09 testing-access clarification qualifies the P237 yes"]
                    self.assertEqual(result[0], expected, result[1])

    @unittest.skipUnless(DEFAULT_FINDINGS.is_file(), "Private saved model responses are not available.")
    def test_replayed_qualified_collisions_restore_saved_proposals_without_rewriting_history(self):
        """Replay the main run's one collision and, when available, the pilot's eight; no model calls or writes."""
        sources = [(DEFAULT_FINDINGS, "1.2.0", 1)]
        pilot = DEFAULT_FINDINGS.with_name("pilot-1.3.0.sqlite")
        if pilot.is_file():
            sources.append((pilot, "1.3.0", 8))
        physical = dict(self.connection.execute(
            "SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        for path, prompt_version, expected_collisions in sources:
            with self.subTest(database=path.name):
                original = path.read_bytes()
                store = open_findings(path, readonly=True)
                self.addCleanup(store.close)
                restored = 0
                for run_record in load_runs(store):
                    if run_record.prompt_version != prompt_version:
                        continue
                    for outcome in run_record.batches:
                        if not any("one role" in reason for reason in outcome.rejected):
                            continue
                        proposal = Proposal.model_validate_json(outcome.response)
                        _, findings, rejected, _ = realize(
                            self.connection, store, self.batches[outcome.batch_id], proposal,
                            run_record.source_versions[outcome.batch_id.split(":")[0]],
                            set(outcome.input_passage_ids), self.expected)
                        self.assertFalse(any("one role" in reason for reason in rejected), outcome.batch_id)
                        for proposed in proposal.findings:
                            if not set(proposed.supporting) & set(proposed.qualifying):
                                continue
                            self.assertTrue(any(text.strip() for text in proposed.qualifications))
                            finding = next(f for f in findings if f.statement == proposed.statement)
                            self.assertEqual(finding.supporting, [physical[c] for c in proposed.supporting])
                            self.assertEqual(finding.qualifying,
                                             [physical[c] for c in proposed.qualifying if c not in proposed.supporting])
                            self.assertEqual(finding.qualifications, proposed.qualifications)
                            self.assertEqual(finding.review_state, "proposed")
                            restored += 1
                self.assertEqual(restored, expected_collisions)
                self.assertEqual(path.read_bytes(), original)

    def gold_store(self, name, contexts, findings, *, batches=None):
        """A store for gold evaluation: an analyst seed, or a model run whose section outcomes are `batches`."""
        versions = {d.id: SourceVersion(source_sha256=d.source_sha256, extraction_sha256=d.extraction_sha256)
                    for d in self.corpus.documents}
        if batches is None:
            run = ExtractionRun(id=name, model="none", prompt_version="analyst-seed", source_versions=versions)
        else:
            outcomes = [BatchOutcome(batch_id=batch_id, status=status, input_passage_ids=[],
                                     input_fingerprint=FP, error="HTTP 503" if status == "failed" else None,
                                     finding_ids=[f.id for f in findings if status == "findings"
                                                  and f.document_id == batch_id.split(":")[0]][:1])
                        for batch_id, status in batches.items()]
            run = ExtractionRun(id=name, model="test-model", prompt_version="test", source_versions=versions,
                                batches=outcomes)
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        store = open_findings(Path(directory.name) / f"{name}.sqlite")
        self.addCleanup(store.close)
        save_run(store, run, contexts, findings)
        return {case: (status, detail) for case, status, detail in evaluate(store, self.connection)}

    def gold_records(self):
        """Contexts for the three interviews, and a finding factory keyed by citation IDs."""
        documents = {d.id: d for d in self.corpus.documents}
        physical = dict(self.connection.execute("SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))

        def version(document_id):
            return dict(source_sha256=documents[document_id].source_sha256,
                        extraction_sha256=documents[document_id].extraction_sha256)

        contexts = [
            CompanyContext(id="E1:C01", document_id="E1", company_name="Xena", aliases=["Abzena"],
                           employment="current", supporting=[physical["E1:P016"]], origin="model", **version("E1")),
            CompanyContext(id="E1:C02", document_id="E1", company_name="Thermo Fisher", employment="historical",
                           supporting=[physical["E1:P021"]], origin="model", **version("E1")),
            CompanyContext(id="E2:C01", document_id="E2", company_name="Calloway Biosciences", employment="current",
                           supporting=[physical["E2:P015"]], origin="model", **version("E2")),
            CompanyContext(id="E3:C01", document_id="E3", company_name="Solara Renewables", employment="current",
                           supporting=[physical["E3:P017"]], origin="model", **version("E3")),
            CompanyContext(id="E3:C02", document_id="E3", company_name="Industrial Services Company",
                           employment="historical", supporting=[physical["E3:P017"]], origin="model",
                           **version("E3"))]
        numbers: dict[str, int] = {}

        def finding(context_id, supporting, statement, **extra):
            document = context_id.split(":")[0]
            numbers[document] = numbers.get(document, 0) + 1
            for role in ("qualifying", "context"):
                if role in extra:
                    extra[role] = [physical[c] for c in extra[role]]
            return Finding(**{**dict(id=f"{document}:F{numbers[document]:03d}", document_id=document,
                                     context_id=context_id, kind="implementation", statement=statement,
                                     evidence_type="firsthand_report", supporting=[physical[c] for c in supporting],
                                     origin="model", **version(document)), **extra})

        return contexts, finding

    def test_gold_cases_whose_sections_did_not_complete_report_not_run(self):
        contexts, finding = self.gold_records()
        thermo = finding("E1:C02", ["E1:P021"], "Thermo Fisher ran ServiceNow and had started implementing GRC.",
                         kind="vendor_relationship", vendor="ServiceNow", relationship="deployed")
        partial = self.gold_store("partial", contexts, [thermo], batches={"E1:S02": "findings", "E1:S03": "failed"})
        for case in ("G15 Thermo staffing is not full-time headcount",  # E1:S05 never ran
                     "G10 Thermo ranking and decisive reason remain distinct",  # E1:S03 failed
                     "G01/G02 current employer and Thermo Fisher stay distinct",  # E1:S01 never ran
                     "G14 license prices retain company, amount, unit, currency, and stated basis",
                     "G07/G08 interviewer premises are not supporting evidence",  # needs every section
                     "G18 no market-share percentage is inferred from interviews"):
            with self.subTest(case=case):
                self.assertEqual(partial[case][0], "not_run", partial[case][1])
        physical = dict(self.connection.execute("SELECT citation_id, id FROM passages WHERE citation_id IS NOT NULL"))
        staffing_section = next(b.id for b in self.batches.values() if physical["E1:P171"] in b.passage_ids)
        self.assertIn(f"sections not processed: {staffing_section}",
                      partial["G15 Thermo staffing is not full-time headcount"][1])
        self.assertNotIn("missing E1", partial["G14 license prices retain company, amount, unit, currency, and "
                                               "stated basis"][1])
        complete = self.gold_store("complete", contexts, [thermo], batches={
            batch.id: "findings" if batch.id == "E1:S02" else "no_findings" for batch in model_batches(
                list(self.batches.values()))})
        self.assertEqual(complete["G15 Thermo staffing is not full-time headcount"][0], "absent")
        self.assertEqual(complete["G07/G08 interviewer premises are not supporting evidence"][0], "pass")
        bmc = complete["G13 BMC timing tension remains linked and unreconciled"]
        self.assertEqual(bmc[0], "not_run")  # the two timing accounts sit in different sections; linking never ran
        self.assertIn("E2:L01", bmc[1])

    def test_gold_accepts_correct_wording_it_used_to_flag(self):
        contexts, finding = self.gold_records()
        aiops = finding("E2:C01", ["E2:P022"], "Calloway Biosciences is evaluating the full AIOps suite and has not "
                        "turned it on yet.", kind="vendor_relationship", vendor="BMC Helix", relationship="evaluated")
        staffing = [
            finding("E1:C02", ["E1:P171"], "About 3-5 employees maintained ServiceNow at Thermo Fisher.",
                    quantity=Quantity(raw="about 3-5 employees", low=3, high=5, approximate=True,
                                      unit="employees", basis="to maintain the system"),
                    qualifications=["These are part-time allocations, not full-time equivalents."]),
            finding("E1:C02", ["E1:P171"], "Each spent around 30-40 percent of their time on it.",
                    quantity=Quantity(raw="around 30-40%", low=30, high=40, approximate=True, unit="percent")),
            finding("E1:C02", ["E1:P171"], "Partners were engaged for major upgrades or complex flows.")]
        ivanti = finding("E3:C02", ["E3:P017"], "The larger industrial services company ran Ivanti Service Manager, "
                         "formerly Ivanti Heat, on-prem.", kind="vendor_relationship",
                         vendor="Ivanti Service Manager", relationship="deployed")
        neurons = finding("E2:C01", ["E2:P039"], "Calloway evaluated Ivanti Neurons for ITSM.",
                          kind="vendor_relationship", vendor="Ivanti Neurons", relationship="evaluated")
        returned = finding("E1:C02", ["E1:P057"], "ServiceNow gave Thermo Fisher validated GxP change workflows.",
                           context=["E1:P055"])
        reporting = finding("E3:C01", ["E3:P127"], "Solara builds a SOC 2 audit report manually in Excel.",
                            qualifying=["E3:P131"], qualifications=[
                                "Native reporting does not flex without the analytics add-on, which Solara has not "
                                "purchased."])
        results = self.gold_store("wording", contexts, [aiops, *staffing, ivanti, neurons, returned, reporting])
        modules = results["G05 module status stays installed/evaluated/considered/unpurchased"][1]
        self.assertNotIn("AIOps", modules)
        self.assertNotIn("analytics add-on", modules)
        for case in ("G15 Thermo staffing is not full-time headcount",
                     "G17 Ivanti products and experience are not conflated",
                     "G02 compliance discussion returns to Thermo Fisher"):
            with self.subTest(case=case):
                self.assertEqual(results[case][0], "pass", results[case][1])

    def test_gold_still_fails_the_errors_behind_the_relaxed_checks(self):
        contexts, finding = self.gold_records()
        aiops = finding("E2:C01", ["E2:P022"], "Calloway Biosciences turned on the full AIOps suite.",
                        kind="vendor_relationship", vendor="BMC Helix", relationship="deployed")
        fte = finding("E1:C02", ["E1:P171"], "About 3-5 full-time employees maintained ServiceNow, with partners.",
                      quantity=Quantity(raw="about 3-5 employees", low=3, high=5, basis="employees"))
        share = finding("E1:C02", ["E1:P171"], "Each spent 30-40 percent of their time.",
                        quantity=Quantity(raw="30-40%", low=30, high=40, unit="percent"))
        ivanti = finding("E3:C01", ["E3:P017"], "Solara ran Ivanti Service Manager, formerly Ivanti Heat, on-prem.",
                         kind="vendor_relationship", vendor="Ivanti Service Manager", relationship="deployed")
        returned = finding("E1:C02", ["E1:P057"], "ServiceNow gave Thermo Fisher validated GxP change workflows.")
        results = self.gold_store("errors", contexts, [aiops, fte, share, ivanti, returned])
        self.assertIn("AIOps", results["G05 module status stays installed/evaluated/considered/unpurchased"][1])
        for case in ("G15 Thermo staffing is not full-time headcount",
                     "G17 Ivanti products and experience are not conflated",
                     "G02 compliance discussion returns to Thermo Fisher"):
            with self.subTest(case=case):
                self.assertEqual(results[case][0], "fail", results[case][1])

    def judge_fixture(self, findings, *, seed="judge", documents=("E1", "E2", "E3")):
        """An analyst-seeded store holding `findings` plus the gold contexts of `documents`, for judge tests."""
        contexts = [c for c in self.gold_records()[0] if c.document_id in documents]
        versions = {d.id: SourceVersion(source_sha256=d.source_sha256, extraction_sha256=d.extraction_sha256)
                    for d in self.corpus.documents}
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        store = open_findings(Path(directory.name) / f"{seed}.sqlite")
        self.addCleanup(store.close)
        save_run(store, ExtractionRun(id=seed, model="none", prompt_version="analyst-seed", source_versions=versions),
                 contexts, findings)
        return store, Path(directory.name) / "judgments.jsonl"

    def testing_findings(self, finding=None):
        finding = finding or self.gold_records()[1]
        return [finding("E1:C02", ["E1:P241"], "Unexpected testing cost came from underestimating how many business "
                        "users needed access to test workflows.", qualifying=["E1:P237", "E1:P245"],
                        qualifications=["The initial yes was clarified: testing access, rather than purchasing "
                                        "separate non-production environments."])]

    def verdicts(self, **items):
        return json.dumps({"verdicts": [dict(item=item, verdict=verdict, finding_ids=ids, reason="As stated.")
                                        for item, (verdict, ids) in items.items()]})

    def test_rubrics_read_only_their_case_passages_and_never_reach_the_extractor(self):
        for case, items in extract.GOLD_RUBRICS.items():
            self.assertIn(case, extract.CASE_PASSAGES)
            self.assertEqual(len({item.id for item in items}), len(items))
            for item in items:
                with self.subTest(case=case, item=item.id):
                    self.assertTrue(item.passages)
                    self.assertLessEqual(set(item.passages), set(extract.CASE_PASSAGES[case]))
                    for prompt in (extract.PROMPT, extract.LINK_PROMPT):
                        self.assertNotIn(item.text, prompt)
        self.assertNotIn("G04", extract.GOLD_RUBRICS)  # cases that stay fully deterministic
        self.assertNotIn("G14", extract.GOLD_RUBRICS)

    def test_judge_input_is_blind_canonical_and_limited_to_code_chosen_candidates(self):
        _, finding = self.gold_records()
        testing = self.testing_findings(finding)
        unrelated = finding("E1:C02", ["E1:P171"], "About 3-5 employees maintained ServiceNow.")
        store, cache = self.judge_fixture([*testing, unrelated])
        fake = FakeModel([self.verdicts(M1=("met", ["E1:F001"]), M2=("met", ["E1:F001"]), N1=("clear", []))])
        judge_cases(store, self.connection, call=fake, cache=cache, model="judge-model", cases=["G09"])
        system, user = fake.calls[0]
        self.assertEqual(system, extract.JUDGE_PROMPT)
        self.assertIn("E1:F001", user)
        self.assertNotIn("E1:F002", user)  # the staffing finding cites none of G09's passages
        self.assertIn("[E1:P241] Expert 1", user)  # quoted from canonical passages
        self.assertIn("we underestimated the number of business users", user)
        for leak in (str(cache.parent), "judge.sqlite", "not_run", "absent", "deterministic", "gold", "run-1."):
            self.assertNotIn(leak, user)

    def test_judge_replies_are_rejected_unless_complete_and_grounded_in_the_items_passages(self):
        store, cache = self.judge_fixture(self.testing_findings())
        bad = {
            "unknown finding": self.verdicts(M1=("met", ["E1:F999"]), M2=("met", ["E1:F001"]), N1=("clear", [])),
            "missing item": self.verdicts(M1=("met", ["E1:F001"]), N1=("clear", [])),
            "wrong verdict kind": self.verdicts(M1=("clear", []), M2=("met", ["E1:F001"]), N1=("clear", [])),
            "met without evidence": self.verdicts(M1=("met", []), M2=("met", ["E1:F001"]), N1=("clear", [])),
            "not json": "The findings look fine.",
        }
        for name, reply in bad.items():
            with self.subTest(name):
                result, = judge_cases(store, self.connection, call=FakeModel([reply]), cache=cache.with_name(
                    f"{name}.jsonl"), model="judge-model", cases=["G09"])
                self.assertEqual(result[1], "judge_unavailable", result[2])
        # G13's link item needs a finding that cites both timing passages, not one of them.
        _, finding = self.gold_records()
        seven = finding("E2:C01", ["E2:P089"], "The rollout took about seven months, faster than expected.")
        store, cache = self.judge_fixture([seven], seed="bmc")
        reply = self.verdicts(M1=("met", ["E2:F001"]), M2=("met", ["E2:F001"]), N1=("clear", []))
        result, = judge_cases(store, self.connection, call=FakeModel([reply]), cache=cache, model="judge-model",
                              cases=["G13"])
        self.assertEqual(result[1], "judge_unavailable")
        self.assertIn("E2:F001", result[2])

    def test_judge_verdicts_combine_so_a_violation_fails_even_when_every_must_is_met(self):
        cases = {"pass": self.verdicts(M1=("met", ["E1:F001"]), M2=("met", ["E1:F001"]), N1=("clear", [])),
                 "fail": self.verdicts(M1=("met", ["E1:F001"]), M2=("not_met", []), N1=("clear", [])),
                 "violated": self.verdicts(M1=("met", ["E1:F001"]), M2=("met", ["E1:F001"]),
                                           N1=("violated", ["E1:F001"]))}
        for expected, reply in cases.items():
            with self.subTest(expected):
                store, cache = self.judge_fixture(self.testing_findings(), seed=expected)
                result, = judge_cases(store, self.connection, call=FakeModel([reply]), cache=cache,
                                      model="judge-model", cases=["G09"])
                self.assertEqual(result[1], "pass" if expected == "pass" else "fail", result[2])

    def test_judge_fails_closed_caches_and_calls_only_for_scoreable_cases(self):
        store, cache = self.judge_fixture(self.testing_findings())
        down, = judge_cases(store, self.connection, call=FakeModel([RuntimeError("HTTP 529 overloaded")]),
                            cache=cache, model="judge-model", cases=["G09"])
        self.assertEqual(down[1], "judge_unavailable")
        reply = self.verdicts(M1=("met", ["E1:F001"]), M2=("met", ["E1:F001"]), N1=("clear", []))
        fake = FakeModel([reply])
        first, = judge_cases(store, self.connection, call=fake, cache=cache, model="judge-model", cases=["G09"])
        again, = judge_cases(store, self.connection, call=fake, cache=cache, model="judge-model", cases=["G09"])
        self.assertEqual((first[1], again[1], len(fake.calls)), ("pass", "pass", 1))  # the unchanged case is cached
        none = FakeModel([])
        (absent,) = judge_cases(store, self.connection, call=none, cache=cache, model="judge-model", cases=["G15"])
        e1_only, _ = self.judge_fixture(self.testing_findings(), seed="e1", documents=("E1",))
        (unrun,) = judge_cases(e1_only, self.connection, call=none, cache=cache, model="judge-model", cases=["G13"])
        self.assertEqual((absent[1], unrun[1], len(none.calls)), ("absent", "not_run", 0))

    def test_claude_replies_other_than_a_finished_answer_are_errors(self):
        finished = {"content": [{"type": "thinking", "thinking": ""}, {"type": "text", "text": "{}"}],
                    "stop_reason": "end_turn", "model": "claude-sonnet-5", "usage": {"output_tokens": 5}}
        self.assertEqual(extract.claude_reply(finished).text, "{}")
        for stop in ("refusal", "max_tokens"):
            with self.subTest(stop), self.assertRaises(RuntimeError):
                extract.claude_reply({**finished, "stop_reason": stop})

    def test_calibration_scores_repeats_against_frozen_labels_and_applies_the_bar_to_held_out_only(self):
        store, cache = self.judge_fixture(self.testing_findings())
        path = Path(tempfile.mkdtemp()) / "labels.json"
        self.addCleanup(lambda: path.unlink(missing_ok=True))
        store_path = store.execute("PRAGMA database_list").fetchone()[2]
        entry = dict(store=store_path, case="G09", raw=False, note="clarified correctly")
        path.write_text(json.dumps({"rubric_version": extract.RUBRIC_VERSION, "labels": [
            {**entry, "split": "dev", "expected": "pass"},
            {**entry, "split": "holdout", "expected": "fail", "items": ["N1"], "seeded": True}]}))
        ok = self.verdicts(M1=("met", ["E1:F001"]), M2=("met", ["E1:F001"]), N1=("clear", []))
        fake = FakeModel([ok] * 3)  # identical input: every label reuses the three repeats
        report = calibrate(path, self.connection, call=fake, cache=cache, model="judge-model", repeats=3)
        self.assertEqual(len(fake.calls), 3)
        self.assertEqual(report["dev"]["agreement"], 1.0)
        self.assertEqual(report["holdout"]["false_pass"], 1)  # the seeded error passed: the bar fails
        self.assertEqual(report["holdout"]["seeded_caught"], 0.0)
        self.assertFalse(report["bar_met"])
        path.write_text(json.dumps({"rubric_version": "0.9.0", "labels": []}))
        with self.assertRaises(ValueError):  # labels written for another rubric are refused
            calibrate(path, self.connection, call=fake, cache=cache, model="judge-model")

if __name__ == "__main__":
    unittest.main()
