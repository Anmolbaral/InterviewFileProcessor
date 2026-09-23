"""Parse the supplied DOCX interviews into source-addressable, literal records in one SQLite database."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import zipfile
from contextlib import closing
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Literal
from xml.etree import ElementTree as ET

import docx
from lxml import etree
from pydantic import BaseModel, ConfigDict, Field, TypeAdapter, model_validator

BASE = Path(__file__).resolve().parent
DEFAULT_MANIFEST = BASE / "analysis/transcripts/manifest.json"
DEFAULT_DATABASE = BASE / "data/parsed/transcripts.sqlite"
PARSER_VERSION = "1.1.0"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TURN = re.compile(r"(?P<speaker>Expert\s+\d+|AI Interviewer)\s+(?P<time>\d{2}:[0-5]\d:[0-5]\d)")
ANCILLARY = re.compile(r"word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml")
ANOMALY = re.compile(r"\[AUDIO[^\]\r\n]*\]")
BODY_MARKERS = {W + "bookmarkStart", W + "bookmarkEnd", W + "proofErr"}  # positions, not content
INTERVIEWER = "AI Interviewer"  # ponytail: the two-speaker convention of these transcripts
CHUNKING_VERSION = "1.0.0"  # covers the grouping rule and the header rule; bump when either changes
JSON_COLUMNS = ("warnings", "ancillary_parts", "anomaly_markers", "speaker_labels", "header_passage_ids")
SCHEMA = """
CREATE TABLE IF NOT EXISTS documents (
    id TEXT PRIMARY KEY,
    position INTEGER NOT NULL,
    source_filename TEXT NOT NULL,
    source_sha256 TEXT NOT NULL,
    parser_version TEXT NOT NULL,
    python_docx_version TEXT NOT NULL,
    extraction_sha256 TEXT NOT NULL,
    coverage TEXT NOT NULL,
    warnings TEXT NOT NULL,
    ancillary_parts TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS passages (
    id TEXT PRIMARY KEY,
    document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    paragraph_index INTEGER NOT NULL,
    citation_id TEXT UNIQUE,
    text TEXT NOT NULL,
    kind TEXT NOT NULL,
    style_name TEXT NOT NULL,
    heading_basis TEXT,
    section_id TEXT REFERENCES passages(id),
    turn_id TEXT REFERENCES passages(id),
    speaker_label TEXT,
    timestamp_raw TEXT,
    anomaly_markers TEXT NOT NULL,
    raw_xml TEXT NOT NULL,
    UNIQUE (document_id, paragraph_index)
);
CREATE TABLE IF NOT EXISTS chunks (
    id TEXT PRIMARY KEY,
    document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    ordinal INTEGER NOT NULL,
    level TEXT NOT NULL CHECK (level IN ('exchange', 'unattributed')),
    section_id TEXT REFERENCES passages(id),
    speaker_labels TEXT NOT NULL,
    text TEXT NOT NULL,
    header TEXT NOT NULL,
    header_passage_ids TEXT NOT NULL,
    chunking_version TEXT NOT NULL,
    UNIQUE (document_id, ordinal)
);
CREATE TABLE IF NOT EXISTS chunk_passages (
    chunk_id TEXT NOT NULL REFERENCES chunks(id) ON DELETE CASCADE,
    passage_id TEXT NOT NULL UNIQUE REFERENCES passages(id) ON DELETE CASCADE,
    position INTEGER NOT NULL,
    PRIMARY KEY (chunk_id, position)
);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
    chunk_id UNINDEXED, text, header, tokenize = 'porter unicode61 remove_diacritics 2'
);
"""


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"


def fingerprint(value: object) -> str:
    return hashlib.sha256(json_text(value).encode("utf-8")).hexdigest()


class Record(BaseModel):
    model_config = ConfigDict(strict=True, extra="forbid")


class Source(Record):
    # The reference manifest also contains exploratory word counts and notes.
    model_config = ConfigDict(strict=True, extra="ignore")
    expert: int = Field(ge=1)
    source: str
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    paragraphs: int | None = Field(default=None, ge=0)
    extraction: str | None = None


class Passage(Record):
    id: str
    citation_id: str | None
    paragraph_index: int = Field(ge=1)
    text: str
    kind: Literal["blank", "heading", "turn_marker", "text"]
    style_name: str
    heading_basis: Literal["style", "bold_underline"] | None = None
    section_id: str | None = None
    turn_id: str | None = None
    speaker_label: str | None = None
    timestamp_raw: str | None = None
    anomaly_markers: list[str] = Field(default_factory=list)
    raw_xml: str  # the original paragraph element, so unmodeled formatting is never lost


class AncillaryPart(Record):
    part_name: str
    paragraphs: list[str]
    field_instructions: list[str]


class ParsedDocument(Record):
    id: str = Field(pattern=r"^E[1-9]\d*$")
    source_filename: str
    source_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    parser_version: str
    python_docx_version: str
    extraction_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    coverage: Literal["supported_text", "partial"]
    passages: list[Passage]
    ancillary_parts: list[AncillaryPart]
    warnings: list[str]

    @model_validator(mode="after")
    def check_references_and_fingerprint(self) -> ParsedDocument:
        seen: dict[str, Passage] = {}
        citation_number = 0
        for index, passage in enumerate(self.passages, 1):
            if passage.id != f"{self.id}:B{index:04d}" or passage.paragraph_index != index:
                raise ValueError("Passage IDs/locations must follow original paragraph order")
            citation_number += bool(passage.text.strip())
            expected = f"{self.id}:P{citation_number:03d}" if passage.text.strip() else None
            if passage.citation_id != expected:
                raise ValueError("Legacy citation IDs must follow nonempty paragraph order")
            if (passage.kind == "blank") != (not passage.text.strip()):
                raise ValueError("Blank classification must match preserved text")
            seen[passage.id] = passage
            for ref, kind in [(passage.section_id, "heading"), (passage.turn_id, "turn_marker")]:
                if ref is not None and (ref not in seen or seen[ref].kind != kind):
                    raise ValueError("Context reference must resolve to a preceding source marker")
            if passage.kind in ("heading", "blank") and passage.turn_id is not None:
                raise ValueError("Headings and blank paragraphs are not speech")
            if passage.turn_id is not None:
                marker = seen[passage.turn_id]
                match = TURN.fullmatch(marker.text.strip())
                if not match or (passage.speaker_label, passage.timestamp_raw) != match.groups():
                    raise ValueError("Speaker/time metadata must match the cited literal turn label")
            elif passage.speaker_label is not None or passage.timestamp_raw is not None:
                raise ValueError("Speaker/time metadata requires a source turn marker")
        if self.coverage == "partial" and not self.warnings:
            raise ValueError("Partial coverage must include an explanation")
        if fingerprint(self.model_dump(exclude={"extraction_sha256"})) != self.extraction_sha256:
            raise ValueError("Extraction fingerprint mismatch: regenerate/revalidate this data")
        return self


class Corpus(Record):
    schema_version: Literal["1"] = "1"
    documents: list[ParsedDocument] = Field(min_length=1)

    @model_validator(mode="after")
    def unique_documents(self) -> Corpus:
        ids = [document.id for document in self.documents]
        if len(ids) != len(set(ids)):
            raise ValueError("Each expert/source ID must be unique")
        return self


def source_path(value: str, manifest_dir: Path) -> Path:
    path = Path(value).expanduser()
    return (path if path.is_absolute() else manifest_dir / path).resolve()


def xml_paragraph_text(paragraph: ET.Element) -> str:
    """Read run content wherever it sits inside the paragraph, never paragraph-formatting tab stops."""
    pieces = []
    for run in paragraph.iter(W + "r"):
        for element in run:
            if element.tag == W + "t":
                pieces.append(element.text or "")
            elif element.tag == W + "tab":
                pieces.append("\t")
            elif element.tag == W + "cr" or (
                element.tag == W + "br" and element.get(W + "type", "textWrapping") == "textWrapping"
            ):
                pieces.append("\n")
    return "".join(pieces)


def paragraph_xml(paragraph: etree._Element) -> str:
    """The original paragraph element with only the namespace declarations it actually uses."""
    copied = deepcopy(paragraph)
    etree.cleanup_namespaces(copied)
    return etree.tostring(copied, encoding="unicode", with_tail=False)


def inspect_parts(payload: bytes) -> tuple[list[AncillaryPart], list[str], bool]:
    """Small coverage check, not a general DOCX conversion engine."""
    ancillary, warnings = [], []
    partial = False
    unsupported = {W + tag: tag for tag in (
        "tbl", "txbxContent", "ins", "del", "drawing", "pict", "object", "sdt", "altChunk", "sym",
    )}
    unsupported[W + "tbl"] = "table"
    unsupported["{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"] = "equation"
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        names = ["word/document.xml"] + sorted(name for name in archive.namelist() if ANCILLARY.fullmatch(name))
        for name in names:
            root = ET.fromstring(archive.read(name))
            found = sorted({unsupported[element.tag] for element in root.iter() if element.tag in unsupported})
            if name == "word/document.xml":
                body = root.find(W + "body")
                if body is None:
                    raise ValueError("DOCX has no document body")
                found += sorted({element.tag.rsplit("}", 1)[-1] for element in body
                                 if element.tag not in (W + "p", W + "sectPr")
                                 and element.tag not in unsupported and element.tag not in BODY_MARKERS})
            if found:
                partial = True
                warnings.append(f"{name}: unparsed/partially parsed structures: {', '.join(found)}")
            fields = [element.text or "" for element in root.iter(W + "instrText")]
            fields += [element.get(W + "instr", "") for element in root.iter(W + "fldSimple")]
            if fields:
                warnings.append(f"{name}: field values are cached text, not recalculated (e.g. page numbers).")
            if name != "word/document.xml":
                ancillary.append(AncillaryPart(
                    part_name=name,
                    paragraphs=[xml_paragraph_text(p) for p in root.iter(W + "p")],
                    field_instructions=fields,
                ))
    return ancillary, warnings, partial


def heading_basis(paragraph: docx.text.paragraph.Paragraph) -> str | None:
    if re.match(r"^(Heading \d+|Title|Subtitle)$", paragraph.style.name):
        return "style"
    runs = [run for run in paragraph.runs if run.text.strip()]
    # ponytail: these interviews use bold+underline headings; new layouts need review.
    if runs and all(run.bold and run.underline for run in runs):
        return "bold_underline"
    return None


def read_source(source: Source, manifest_dir: Path) -> tuple[Path, bytes]:
    """Read the original file and confirm its content hash is the one recorded in the manifest."""
    path = source_path(source.source, manifest_dir)
    payload = path.read_bytes()
    if hashlib.sha256(payload).hexdigest() != source.sha256:
        raise ValueError(f"{path.name}: source hash mismatch; review the changed original before updating the manifest")
    return path, payload


def parse_source(source: Source, manifest_dir: Path) -> ParsedDocument:
    path, payload = read_source(source, manifest_dir)
    try:
        ancillary, warnings, partial = inspect_parts(payload)
        document = docx.Document(BytesIO(payload))
    except (zipfile.BadZipFile, ET.ParseError, etree.XMLSyntaxError, KeyError) as error:
        raise ValueError(f"{path.name}: invalid DOCX package: {error}") from error
    document_id = f"E{source.expert}"
    passages: list[Passage] = []
    citation_number = 0
    section_id = None
    turn_id = speaker_label = timestamp_raw = None
    for index, paragraph in enumerate(document.paragraphs, 1):
        passage_id = f"{document_id}:B{index:04d}"
        text = xml_paragraph_text(paragraph._p)  # Complete run content; never strip or normalize it.
        if text != paragraph.text:  # Two independent readers disagree: keep the complete one, flag it.
            partial = True
            warnings.append(f"{passage_id}: run content differs from the python-docx paragraph text; "
                            "stored the run content, review the structure inside this paragraph.")
        nonempty = bool(text.strip())
        citation_number += nonempty
        match = TURN.fullmatch(text.strip())
        basis = heading_basis(paragraph) if nonempty and not match else None
        kind = "blank" if not nonempty else "turn_marker" if match else "heading" if basis else "text"
        if kind == "heading":
            section_id = passage_id
            turn_id = speaker_label = timestamp_raw = None
        elif match:
            turn_id = passage_id
            speaker_label, timestamp_raw = match.groups()
        elif re.fullmatch(r"(?:Expert\s+\d+|AI Interviewer)\s+\S+:\S+", text.strip()):
            turn_id = speaker_label = timestamp_raw = None
            warnings.append(f"{passage_id}: unrecognized turn timestamp; speaker context left unknown.")
        is_speech = kind in ("turn_marker", "text")
        passages.append(Passage(
            id=passage_id,
            citation_id=f"{document_id}:P{citation_number:03d}" if nonempty else None,
            paragraph_index=index, text=text, kind=kind, style_name=paragraph.style.name,
            heading_basis=basis, section_id=section_id,
            turn_id=turn_id if is_speech else None,
            speaker_label=speaker_label if is_speech else None,
            timestamp_raw=timestamp_raw if is_speech else None,
            anomaly_markers=ANOMALY.findall(text),
            raw_xml=paragraph_xml(paragraph._p),
        ))
    if not any(passage.kind == "turn_marker" for passage in passages):
        warnings.append("no speaker turn markers recognized; stored speech is unattributed.")
    if source.paragraphs is not None and citation_number != source.paragraphs:
        raise ValueError(f"{path.name}: nonempty paragraph count changed; verify the existing citation mapping")
    record = dict(
        id=document_id, source_filename=path.name, source_sha256=source.sha256,
        parser_version=PARSER_VERSION, python_docx_version=docx.__version__,
        coverage="partial" if partial else "supported_text",
        passages=[passage.model_dump() for passage in passages],
        ancillary_parts=[part.model_dump() for part in ancillary], warnings=warnings,
    )
    return ParsedDocument.model_validate({**record, "extraction_sha256": fingerprint(record)})


def read_manifest(path: Path) -> list[Source]:
    sources = TypeAdapter(list[Source]).validate_python(json.loads(path.read_text(encoding="utf-8")))
    ids = [source.expert for source in sources]
    if not sources or len(ids) != len(set(ids)):
        raise ValueError("Manifest must contain sources with unique expert IDs")
    return sources


def parse_manifest(path: Path) -> Corpus:
    path = path.resolve()
    return Corpus(documents=[parse_source(source, path.parent) for source in read_manifest(path)])


def open_database(path: Path, *, readonly: bool = False) -> sqlite3.Connection:
    path = path.resolve()
    connection = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True) if readonly else sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    if not any(row[0] == "ENABLE_FTS5" for row in connection.execute("PRAGMA compile_options")):
        connection.close()
        raise ValueError("this Python's sqlite3 lacks FTS5, which the keyword index needs; "
                         "use a build whose PRAGMA compile_options lists ENABLE_FTS5")
    return connection


def insert_rows(connection: sqlite3.Connection, table: str, rows: list[dict]) -> None:
    if rows:
        keys = list(rows[0])
        connection.executemany(
            f"INSERT INTO {table} ({', '.join(keys)}) VALUES ({', '.join('?' * len(keys))})",
            [[json.dumps(row[key], ensure_ascii=False) if key in JSON_COLUMNS else row[key] for key in keys]
             for row in rows],
        )


def save_corpus(connection: sqlite3.Connection, corpus: dict, replace: set[str] | None = None) -> None:
    """Store the corpus in one transaction. Only documents in `replace` (default: all) are rewritten;
    the others keep their rows. Documents no longer in the corpus are removed. Derived chunks and the
    keyword index are rebuilt in the same transaction."""
    connection.executescript(SCHEMA)
    ids = [document["id"] for document in corpus["documents"]]
    with connection:
        connection.execute(f"PRAGMA user_version = {int(corpus['schema_version'])}")
        connection.execute(f"DELETE FROM documents WHERE id NOT IN ({', '.join('?' * len(ids))})", ids)
        for position, document in enumerate(corpus["documents"]):
            if replace is not None and document["id"] not in replace:
                connection.execute("UPDATE documents SET position = ? WHERE id = ?", (position, document["id"]))
                continue
            connection.execute("DELETE FROM documents WHERE id = ?", (document["id"],))
            insert_rows(connection, "documents",
                        [{**{key: value for key, value in document.items() if key != "passages"}, "position": position}])
            insert_rows(connection, "passages",
                        [{**passage, "document_id": document["id"]} for passage in document["passages"]])
        save_chunks(connection, corpus["documents"])


def read_rows(connection: sqlite3.Connection, query: str, *parameters: object) -> list[dict]:
    return [{key: json.loads(row[key]) if key in JSON_COLUMNS else row[key] for key in row.keys()}
            for row in connection.execute(query, parameters)]


def load_documents(connection: sqlite3.Connection, document_id: str | None = None) -> list[dict]:
    documents = (read_rows(connection, "SELECT * FROM documents WHERE id = ?", document_id) if document_id
                 else read_rows(connection, "SELECT * FROM documents ORDER BY position"))
    for document in documents:
        del document["position"]
        document["passages"] = read_rows(
            connection, "SELECT * FROM passages WHERE document_id = ? ORDER BY paragraph_index", document["id"])
        for passage in document["passages"]:
            del passage["document_id"]
    return documents


def load_corpus(connection: sqlite3.Connection) -> dict:
    version = connection.execute("PRAGMA user_version").fetchone()[0]
    return {"schema_version": str(version), "documents": load_documents(connection)}


def chunk_document(document: dict) -> list[dict]:
    """Group text passages into exchanges: one interviewer turn plus the replies that follow it until the next
    interviewer turn or heading. Text outside any turn forms 'unattributed' chunks. Chunk text and header exist
    for keyword search only; attribute and quote from passages.text through the membership rows."""
    by_id = {passage["id"]: passage for passage in document["passages"]}
    groups: list[list[dict]] = []
    previous = None
    for passage in document["passages"]:
        if passage["kind"] == "heading":
            previous = None
        elif passage["kind"] == "text":
            changed = previous is None or passage["turn_id"] != previous["turn_id"]
            if changed and (previous is None or passage["speaker_label"] == INTERVIEWER
                            or (passage["turn_id"] is None) != (previous["turn_id"] is None)):
                groups.append([])
            groups[-1].append(passage)
            previous = passage
    # ponytail: no size cap; the largest exchange here is 283 words. Split downstream by passage IDs if needed.
    profile = [p for group in groups for p in group if p["turn_id"] is None and p["section_id"] is None]
    chunks = []
    for ordinal, members in enumerate(groups, 1):
        heading = by_id[members[0]["section_id"]] if members[0]["section_id"] else None
        parts = [document["id"], " ".join(" ".join(p["text"].split()) for p in profile),
                 heading["text"].strip() if heading else ""]
        chunks.append(dict(
            id=f"{document['id']}:X{ordinal:04d}", document_id=document["id"], ordinal=ordinal,
            level="exchange" if members[0]["turn_id"] else "unattributed", section_id=members[0]["section_id"],
            speaker_labels=list(dict.fromkeys(p["speaker_label"] for p in members if p["speaker_label"])),
            text="\n".join(p["text"] for p in members),
            header=" | ".join(part for part in parts if part),  # interview profile and section, not attribution
            header_passage_ids=[p["id"] for p in profile] + ([heading["id"]] if heading else []),
            chunking_version=CHUNKING_VERSION, passage_ids=[p["id"] for p in members],
        ))
    return chunks


def load_chunks(connection: sqlite3.Connection, document_id: str | None = None) -> list[dict]:
    where = "WHERE c.document_id = ?" if document_id else ""
    chunks = read_rows(connection, f"SELECT c.* FROM chunks c JOIN documents d ON d.id = c.document_id {where} "
                       "ORDER BY d.position, c.ordinal", *((document_id,) if document_id else ()))
    members: dict[str, list[str]] = {}
    for row in connection.execute("SELECT chunk_id, passage_id FROM chunk_passages ORDER BY chunk_id, position"):
        members.setdefault(row["chunk_id"], []).append(row["passage_id"])
    for chunk in chunks:
        chunk["passage_ids"] = members.get(chunk["id"], [])
    return chunks


def index_rows(connection: sqlite3.Connection) -> tuple[list[tuple], list[tuple]]:
    """(keyword index rows, chunk rows) as comparable (id, text, header) tuples."""
    return ([tuple(r) for r in connection.execute("SELECT chunk_id, text, header FROM chunks_fts ORDER BY chunk_id")],
            [tuple(r) for r in connection.execute("SELECT id, text, header FROM chunks ORDER BY id")])


def save_chunks(connection: sqlite3.Connection, documents: list[dict]) -> None:
    """Inside the caller's transaction: rebuild every derived chunk, membership, and keyword-index row."""
    # ponytail: full rebuild of ~160 rows each run; reconcile per document only if the corpus grows a hundredfold.
    connection.execute("DELETE FROM chunks")
    connection.execute("DELETE FROM chunks_fts")
    for document in documents:
        chunks = chunk_document(document)
        insert_rows(connection, "chunks", [{k: v for k, v in c.items() if k != "passage_ids"} for c in chunks])
        insert_rows(connection, "chunk_passages", [dict(chunk_id=c["id"], passage_id=p, position=i)
                                                   for c in chunks for i, p in enumerate(c["passage_ids"], 1)])
    connection.execute("INSERT INTO chunks_fts (chunk_id, text, header) SELECT id, text, header FROM chunks")


def render_passages(connection: sqlite3.Connection, passage_ids: list[str], *,
                    expected_fingerprints: dict[str, str] | None = None) -> list[dict]:
    """Attributed evidence in source order, including across sections and documents.
    Nonempty requests require an expected extraction fingerprint for every document; unknown/duplicate IDs,
    missing expectations, and version mismatches fail. Run the read-only integrity check before each evidence
    batch: this function compares stored fingerprints, not the original files or a new hash of stored text.
    Quote through these canonical passages, never chunk text or headers."""
    if len(set(passage_ids)) != len(passage_ids):
        raise ValueError("duplicate passage IDs requested")
    if not passage_ids:
        return []
    rows = read_rows(
        connection,
        "SELECT p.id, p.document_id, p.citation_id, p.speaker_label, p.timestamp_raw, p.text, "
        "d.source_sha256, d.extraction_sha256 FROM passages p JOIN documents d ON d.id = p.document_id "
        f"WHERE p.id IN ({', '.join('?' * len(passage_ids))}) ORDER BY d.position, p.paragraph_index", *passage_ids)
    missing = sorted(set(passage_ids) - {row["id"] for row in rows})
    if missing:
        raise ValueError(f"unknown passage IDs: {', '.join(missing)}")
    for row in rows:
        expected = (expected_fingerprints or {}).get(row["document_id"])
        if expected is None:
            raise ValueError(f"{row['document_id']}: expected extraction fingerprint required for every requested "
                             "document; verify the corpus before reading evidence")
        if expected != row["extraction_sha256"]:
            raise ValueError(f"{row['document_id']}: expected extraction fingerprint differs from the stored one; "
                             "run --check and revalidate dependent facts before using this evidence")
    return rows


def search(connection: sqlite3.Connection, query: str, *, document_id: str | None = None,
           limit: int = 10) -> list[dict]:
    """Keyword search over chunks, BM25-ranked with header words at half weight (a bias, not a guarantee).
    Each hit separates evidence (passage_ids) from search context (header, header_passage_ids).
    A query FTS5 cannot parse, such as one containing $ or ?, is retried as its bare words, all required."""
    sql = """
        SELECT c.id, c.document_id, c.level, c.ordinal, c.section_id, h.text AS section_heading,
               c.speaker_labels, c.header, c.header_passage_ids,
               bm25(chunks_fts, 0.0, 1.0, 0.5) AS rank, snippet(chunks_fts, 1, '[', ']', '…', 20) AS snippet
        FROM chunks_fts JOIN chunks c ON c.id = chunks_fts.chunk_id
        LEFT JOIN passages h ON h.id = c.section_id
        WHERE chunks_fts MATCH ? AND (? IS NULL OR c.document_id = ?)
        ORDER BY rank, c.document_id, c.ordinal LIMIT ?"""
    try:
        hits = read_rows(connection, sql, query, document_id, document_id, limit)
    except sqlite3.OperationalError:  # ponytail: no query parser; quoted bare words are the whole fallback
        words = re.findall(r"\w+", query)
        hits = (read_rows(connection, sql, " ".join(f'"{w}"' for w in words), document_id, document_id, limit)
                if words else [])
    for hit in hits:
        hit["passage_ids"] = [row[0] for row in connection.execute(
            "SELECT passage_id FROM chunk_passages WHERE chunk_id = ? ORDER BY position", (hit["id"],))]
    return hits


def neighbors(connection: sqlite3.Connection, chunk_id: str, *, before: int = 1, after: int = 1) -> list[dict]:
    """The surrounding chunks within the same section, anchor included, each with passage_ids. Widening the
    window never crosses a heading; read another section or the whole transcript from passages directly."""
    anchors = read_rows(connection, "SELECT document_id, section_id, ordinal FROM chunks WHERE id = ?", chunk_id)
    if not anchors:
        return []
    anchor = anchors[0]
    return [chunk for chunk in load_chunks(connection, anchor["document_id"])
            if chunk["section_id"] == anchor["section_id"]
            and anchor["ordinal"] - before <= chunk["ordinal"] <= anchor["ordinal"] + after]


def run(manifest: Path, database: Path, check: bool = False) -> tuple[Corpus, list[str]]:
    """Parse changed or new sources into the database, or verify it. Returns the corpus and the IDs whose
    rows were written; a document with the stored content hash, parser, and library versions is reused."""
    manifest, database = manifest.resolve(), database.resolve()
    sources = read_manifest(manifest)
    protected = {manifest}
    for source in sources:
        protected.add(source_path(source.source, manifest.parent))
        if source.extraction:
            protected.add(source_path(source.extraction, manifest.parent))
    if database in protected or database.is_relative_to((BASE / "sources").resolve()):
        raise ValueError("Database must not overwrite a source, reference manifest, exploratory extract, or sources/")
    if check:
        corpus = Corpus(documents=[parse_source(source, manifest.parent) for source in sources])
        with closing(open_database(database, readonly=True)) as connection:
            stored = Corpus.model_validate(load_corpus(connection))
            try:
                chunks = load_chunks(connection)
                indexed, current = index_rows(connection)
            except sqlite3.OperationalError as error:
                raise ValueError("database has no chunk tables; run the parser first") from error
        if stored.model_dump() != corpus.model_dump():
            raise ValueError("Stored database is stale or different; regenerate and revalidate dependent facts")
        if chunks != [chunk for document in stored.model_dump()["documents"] for chunk in chunk_document(document)]:
            raise ValueError("Stored chunks are stale or tampered; rerun the parser to rebuild them")
        if indexed != current:
            raise ValueError("Keyword index does not match the chunks; rerun the parser to rebuild it")
        return corpus, []
    database.parent.mkdir(parents=True, exist_ok=True)
    with closing(open_database(database)) as connection:
        connection.executescript(SCHEMA)
        stored = {row["id"]: (row["source_sha256"], row["parser_version"], row["python_docx_version"])
                  for row in connection.execute("SELECT id, source_sha256, parser_version, python_docx_version "
                                                "FROM documents")}
        documents, written = [], []
        for source in sources:
            document_id = f"E{source.expert}"
            if stored.get(document_id) == (source.sha256, PARSER_VERSION, docx.__version__):
                read_source(source, manifest.parent)  # The file must still carry its recorded content hash.
                try:
                    documents.append(ParsedDocument.model_validate(load_documents(connection, document_id)[0]))
                except ValueError as error:
                    raise ValueError(f"{document_id}: stored rows fail validation; run --check or delete the "
                                     "database before parsing again") from error
            else:
                documents.append(parse_source(source, manifest.parent))
                written.append(document_id)
        corpus = Corpus(documents=documents)
        save_corpus(connection, corpus.model_dump(), replace=set(written))
    return corpus, written


def main() -> int:
    command = argparse.ArgumentParser(description=__doc__)
    command.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    command.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    command.add_argument("--check", action="store_true",
                         help="Reparse every source and verify the stored database without changing it")
    command.add_argument("--search", metavar="QUERY",
                         help="Print keyword-search hits from the stored database, read-only, and exit")
    args = command.parse_args()
    try:
        if args.search is not None:
            corpus, _ = run(args.manifest, args.database, check=True)
            expected = {document.id: document.extraction_sha256 for document in corpus.documents}
            with closing(open_database(args.database, readonly=True)) as connection:
                for hit in search(connection, args.search):
                    print(f"{hit['id']}  {hit['section_heading'] or '(no section)'}  rank={hit['rank']:.2f}")
                    print(f"  header: {hit['header']}  <- {hit['header_passage_ids']}")
                    for passage in render_passages(connection, hit["passage_ids"], expected_fingerprints=expected):
                        print(f"  {passage['citation_id']} {passage['speaker_label'] or '-'} "
                              f"{passage['timestamp_raw'] or '-'}: {passage['text'][:80]!r}")
            return 0
        corpus, written = run(args.manifest, args.database, args.check)
        with closing(open_database(args.database, readonly=True)) as connection:
            counts = {row["document_id"]: (row["n"], row["exchanges"]) for row in connection.execute(
                "SELECT document_id, count(*) AS n, sum(level = 'exchange') AS exchanges FROM chunks "
                "GROUP BY document_id")}
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, ET.ParseError, sqlite3.Error) as error:
        command.exit(1, f"Parser failed: {error}\n")
    print(f"{'Verified' if args.check else 'Updated'} {args.database}")
    for document in corpus.documents:
        total, exchanges = counts.get(document.id, (0, 0))
        status = ("verified" if args.check else "parsed and stored" if document.id in written
                  else "unchanged; stored rows reused")
        print(f"{document.id}: {len(document.passages)} body paragraphs, "
              f"{sum(p.citation_id is not None for p in document.passages)} citations; {document.coverage}; "
              f"{total} chunks ({exchanges} exchanges); {status}")
        for warning in document.warnings:
            print(f"  Warning: {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
